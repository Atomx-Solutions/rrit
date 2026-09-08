"""Apply the reviewed, evidence-based 2026-09-07 audit to the original workbook.

Requires openpyxl. No application data, credentials, task IDs or estimates are changed.
The input annotations and original workbook backup make every change reviewable.
"""
from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
import json
import math
import re
import shutil
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
NAME = 'AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx'
TARGET = ROOT / NAME
DATE = '2026-09-07'
BACKUP = ROOT / 'docs' / 'backups' / (TARGET.stem + '.before-audit-' + DATE + '.xlsx')
AUDIT = json.loads((ROOT / f'docs/backlog-audit-{DATE}.json').read_text(encoding='utf-8'))
AMENDMENTS = json.loads((ROOT / f'docs/backlog-scope-amendments-{DATE}.json').read_text(encoding='utf-8'))
original_bytes = TARGET.read_bytes()
wb = load_workbook(TARGET)
ws = wb['02_BACKLOG']
if ws.max_column != 24:
    raise RuntimeError('Workbook already audited or structure changed; review before reapplying.')
rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1)}
assert len(rows) == 185 and set(rows) == set(AUDIT)
assert set(AMENDMENTS) <= set(rows)
for entry in AUDIT.values():
    for evidence in entry['paths']:
        assert (ROOT / evidence).exists(), evidence
assert not any(entry['status'] == 'Done' for entry in AUDIT.values())

BACKUP.parent.mkdir(exist_ok=True)
if BACKUP.exists():
    assert BACKUP.read_bytes() == original_bytes, 'Existing backup differs; do not overwrite it.'
else:
    shutil.copy2(TARGET, BACKUP)

labels = {
    'In Review': 'Implementuar — në pranim',
    'In Progress': 'Pjesërisht',
    'Not Started': 'Pa filluar',
    'Blocked': 'Bllokuar',
    'Deferred': 'Shtyrë',
    'Done': 'Përfunduar dhe pranuar',
}
colors = {'In Review': 'DCFCE7', 'In Progress': 'FEF3C7', 'Not Started': 'F1F5F9',
          'Blocked': 'FEE2E2', 'Deferred': 'EDE9FE', 'Done': 'BBF7D0'}
navy = '17324D'
muted = '475569'
line = Side(style='thin', color='E2E8F0')
counts = Counter(entry['status'] for entry in AUDIT.values())
original_formulas = {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == 'f'}
changed_contract_cells = set()

def change_preserving_original(sheet, address, value, reason):
    cell = sheet[address]
    cell.comment = Comment(f'Para auditit {DATE}:\n{cell.value}\n\nArsyeja: {reason}', 'Rrit — audit')
    cell.value = value
    changed_contract_cells.add((sheet.title, address))

for col, title in [(25, 'Vlerësimi i zbatimit'), (26, f'Rishikimi real / çfarë mbetet — {DATE}')]:
    cell = ws.cell(4, col, title)
    cell._style = copy(ws['X4']._style)
    cell.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor=navy)
    cell.alignment = Alignment(wrap_text=True, vertical='center')
ws.column_dimensions['Y'].width = 29
ws.column_dimensions['Z'].width = 116
ws.row_dimensions[4].height = max(ws.row_dimensions[4].height or 30, 42)
ws['L4'].comment = Comment('Status kontraktual: Done kërkon acceptance criteria, prova dhe gates. '
    'In Review: implementimi kryesor i scope-it aktual është në kod, por pranim/teste të plota mbeten. '
    'Shiko Y/Z për gjendjen reale; kjo nuk është përqindje gatishmërie prodhimi.', 'Rrit — audit')
ws['M4'].comment = Comment('Kjo kolonë ruan llojin e provës së KËRKUAR nga plani. '
    'Nuk është deklaratë që prova ka kaluar. Evidenca e auditit është në S dhe Z.', 'Rrit — audit')

for task_id, r in rows.items():
    item = AUDIT[task_id]
    ws.cell(r, 12, item['status'])
    ws.cell(r, 12).fill = PatternFill('solid', fgColor=colors[item['status']])
    ws.cell(r, 12).comment = Comment(f'Status para auditit: Not Started\nAudit: {DATE}\n'
        f'{item["done"]}\nMbetet: {item["remaining"]}', 'Rrit — audit')
    ws.cell(r, 25, labels[item['status']])
    details = (f'{task_id} · {ws.cell(r, 4).value}\n'
               f'Realizuar: {item["done"]}\n'
               f'Mbetet / ndryshimi: {item["remaining"]}\n'
               f'Burime: {"; ".join(item["paths"]) if item["paths"] else "Inventari i kodit / kërkesa e scope-it; nuk u gjet artefakt pranimi."}')
    ws.cell(r, 26, details)
    for col in [25, 26]:
        cell = ws.cell(r, col)
        cell.font = Font(name='Aptos', size=10, color=navy, bold=col == 25)
        cell.fill = PatternFill('solid', fgColor=colors[item['status']] if col == 25 else ('F8FAFC' if r % 2 else 'FFFFFF'))
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        cell.border = Border(bottom=line)
    ws.row_dimensions[r].height = min(240, max(112, sum(max(1, math.ceil(len(part) / 112)) for part in details.splitlines()) * 14 + 12))
    evidence = ws.cell(r, 19, f'Audit {DATE} · {task_id}')
    evidence.hyperlink = f'docs/backlog-audit-{DATE}.json'
    evidence.font = Font(name='Aptos', size=10, color='166534', underline='single')
    if item['status'] == 'Deferred' and ws.cell(r, 22).value == 'GA':
        change_preserving_original(ws, f'V{r}', 'Deferred — posts scope', 'Kërkesa e përdoruesit: vetëm postime; ekip/fushata/analitikë shtyhen, nuk shënohen Done.')
    if task_id in AMENDMENTS:
        for col, key in [('E', 'detail'), ('F', 'acceptance')]:
            change_preserving_original(ws, f'{col}{r}', AMENDMENTS[task_id][key],
                'Përshtatur me kërkesat e shprehura: Nuxt UI, biznes real/PostgreSQL, pa demo, EUR/trial, Meta opsionale. Kriteret e sigurisë/prodhimit mbeten të hapura.')

# Repair the original 184-row freeze and incomplete/overlapping validation ranges.
ws.freeze_panes = 'B5'
ws.sheet_view.pane.topLeftCell = 'Y5'
ws.sheet_view.pane.activePane = 'bottomRight'
ws.sheet_view.selection = [Selection(pane='topRight', activeCell='Y1', sqref='Y1'),
                           Selection(pane='bottomLeft', activeCell='A5', sqref='A5'),
                           Selection(pane='bottomRight', activeCell='Y5', sqref='Y5')]
ws.sheet_view.zoomScale = 85
ws.auto_filter.ref = 'A4:Z189'
ws.print_title_rows = '1:4'
ws.data_validations.dataValidation = []
for formula, region in [
    ('"Not Started,Ready,In Progress,In Review,Blocked,Done,Deferred"', 'L5:L189'),
    ('"P0,P1,P2"', 'H5:H189'), ('"Programmer A,Programmer B"', 'I5:J189'),
]:
    dv = DataValidation(type='list', formula1=formula, allow_blank=False)
    dv.errorTitle = 'Zgjidh një vlerë nga lista'
    dv.error = 'Përdor një vlerë të vlefshme të workbook-ut.'
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(region)
ws.conditional_formatting._cf_rules.clear()
for status, color in colors.items():
    ws.conditional_formatting.add('L5:L189', FormulaRule(formula=[f'$L5="{status}"'], fill=PatternFill('solid', fgColor=color)))
ws.conditional_formatting.add('H5:H189', FormulaRule(formula=['$H5="P0"'], fill=PatternFill('solid', fgColor='FEE2E2')))

# S8/S10 counts were empty; S9 incorrectly contained the whole-plan subtotal.
roadmap = wb['01_ROADMAP']
for r in range(5, 16):
    roadmap[f'M{r}'] = f'=COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$H:$H,"P0",\'02_BACKLOG\'!$L:$L,"<>Done",\'02_BACKLOG\'!$L:$L,"<>Deferred")'
    roadmap[f'N{r}'] = f'=COUNTIF(\'02_BACKLOG\'!$B:$B,$A{r})'
    roadmap[f'O{r}'] = f'=COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$L:$L,"Done")'
    roadmap[f'P{r}'] = (f'=IF(COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$L:$L,"Deferred")=N{r},"Deferred",'
        f'IF(O{r}=N{r},"Done",IF(COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$L:$L,"In Progress")+'
        f'COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$L:$L,"In Review")>0,"In Progress",'
        f'IF(COUNTIFS(\'02_BACKLOG\'!$B:$B,$A{r},\'02_BACKLOG\'!$L:$L,"Blocked")>0,"Blocked","Planned"))))')
    for col in 'MNOP':
        roadmap[f'{col}{r}']._style = copy(roadmap[f'{col}5']._style)
for col in 'MNO':
    roadmap[f'{col}17'] = f'=SUM({col}5:{col}15)'
roadmap['P17'] = 'Overall plan — gates open'
roadmap['A2'] = 'Dates/estimates remain the original plan. Status audited 2026-09-07; S8/S9/S10 counters repaired. In Review is not production acceptance; release gates remain mandatory.'

gate_notes = {
    'RG-01': ('In Progress', 'Nuxt/Nitro/PG kod dhe kontrolle lokale; mungon ADR i stack-ut aktual, Node 24/build aktual, CI me DB dhe immutable deploy.'),
    'RG-02': ('Not Started', 'Owner checks në aplikacion ekzistojnë; RLS/runtime role/context dhe suite për të gjitha tabelat nuk janë implementuar.'),
    'RG-03': ('In Progress', 'Better Auth real dhe origin checks; mungojnë role matrix, sensitive-action audit dhe staging auth suite e plotë.'),
    'RG-04': ('Blocked', 'META_APP_ID/SECRET/LOGIN_CONFIG_ID/REDIRECT_URI mungojnë. Duhet owner Meta App, access review, HTTPS/callbacks dhe non-developer account smoke.'),
    'RG-05': ('In Progress', 'Private JPEG/Sharp/rights/EXIF tests kalojnë; signed direct upload, variants dhe lifecycle tests mungojnë.'),
    'RG-06': ('In Progress', 'Responses captions/schema/verified-fact context në kod; pa OpenAI credentials/evals/immutable provenance/cost budgets.'),
    'RG-07': ('In Progress', 'Manual calendar i provuar pa Meta; publisher në kod. Mungojnë live publish, retry/backoff, reconciler dhe worker fault matrix.'),
    'RG-08': ('In Progress', 'Trial 72h dhe entitlements me prova PG; nuk ka checkout/subscriptions/billing webhooks/reservation ledger.'),
    'RG-09': ('In Progress', 'Privacy UI dhe client JSON export; nuk ka durable export/purge/retention/provider deletion callback dhe legal acceptance.'),
    'RG-10': ('In Progress', 'Audit events/worker heartbeat bazë; pa OpenTelemetry/SLO dashboards/alerts/runbook drills.'),
    'RG-11': ('In Progress', 'Indekse/pool dhe budget script bazë; pa load 3x/EXPLAIN/Web Vitals aktual. Workspace fetch/PUT mbetet i plotë.'),
    'RG-12': ('In Progress', 'Disa auth/CSRF/token/media controls; mungon ASVS/security release review me scans aktuale.'),
    'RG-13': ('Not Started', 'Nuk ka backup/PITR verification, restore exercise ose matje RPO/RTO.'),
    'RG-14': ('In Progress', '2 desktop/mobile calendar E2E green; suite e plotë aktuale, cross-browser/a11y/UAT/prod smoke nuk janë mbyllur.'),
    'RG-15': ('Not Started', 'Strategy/campaign/learning deferred nga scope-i aktual. Gate nuk është Passed/Waived dhe rihapet me këto shërbime.'),
    'RG-16': ('In Progress', 'Local content checks dhe structured captions; mungojnë vertical/golden/injection/claim-provenance/override tests.'),
    'RG-17': ('In Progress', 'Nuxt UI/posts-focused UX dhe manual calendar funksional; fault/a11y/visual/performance/UAT acceptance mbetet.'),
    'RG-18': ('In Progress', 'Daily AI request cap/entitlements/env gating; mungojnë monetary budgets/anomaly/admin/audited flags/COGS replay.'),
    'RG-19': ('Not Started', 'Advanced media janë zgjerim deferred pas-GA. Static-only nuk provon carousel/video/story capability gate.'),
}
gates = wb['04_RELEASE_GATES']
for r in range(5, gates.max_row + 1):
    status, note = gate_notes[gates.cell(r, 1).value]
    gates.cell(r, 7, status)
    gates.cell(r, 8, f'docs/backlog-audit-{DATE}.json').hyperlink = f'docs/backlog-audit-{DATE}.json'
    gates.cell(r, 9, f'{DATE}: {note}').alignment = Alignment(wrap_text=True, vertical='top')
    gates.row_dimensions[r].height = max(95, gates.row_dimensions[r].height or 0)
gates.data_validations.dataValidation = []
dv = DataValidation(type='list', formula1='"Not Started,In Progress,Blocked,Passed,Waived"')
gates.add_data_validation(dv)
dv.add('G5:G23')

decision_rows = [
    ('AUD-2026-01', 'Scope: Kosovo/EUR, posts-only; campaigns/strategy/standalone analytics/team deferred.', 'User instruction', 'Current', 'Product owner', 'Reopen explicitly when the corresponding product area is requested; security/operations gates remain.'),
    ('AUD-2026-02', 'Remove demo runtime. Real PostgreSQL business/account; configurable 72-hour trial with email/browser/social identity claims.', 'User instruction', 'Current', 'Product owner', 'Browser cookie is not hardware proof. Provider credentials and billing remain separate launch prerequisites.'),
    ('AUD-2026-03', 'Meta optional for approval/calendar. Calendar manual by default; automatic publication requires explicit opt-in and provider checks.', 'User instruction', 'Current', 'Product owner', 'Never auto-enable earlier manual calendar entries after connecting Meta. Preserve uncertain-publish guards.'),
    ('AUD-2026-04', 'Local signup may bypass email verification; localhost and 127.0.0.1 aliases at the same port are supported.', 'User instruction', 'Current', 'Product owner', 'Public origins still enforce verification; bypass does not mark emailVerified=true.'),
    ('AUD-2026-05', 'Observed stack: Nitro API, Better Auth, direct PostgreSQL and in-process publisher; original FastAPI/Supabase/outbox baseline differs.', 'Needs ADR review', 'Before GA', 'Engineering owner', 'Reconcile deployment boundaries/toolchain/API contracts/RLS. This audit does not waive enterprise security requirements.'),
]
decisions = wb['07_DECISIONS']
for row in decision_rows:
    decisions.append(row)
    r = decisions.max_row
    for c in decisions[r]:
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.font = Font(name='Aptos', size=10, color=navy)
        c.border = Border(bottom=line)
    decisions.row_dimensions[r].height = 95

readme = wb['00_README']
readme.merge_cells('A44:H44')
readme['A44'] = f'Audit i kodit dhe scope-it — {DATE}'
readme['A44'].font = Font(name='Aptos', size=13, bold=True, color='FFFFFF')
readme['A44'].fill = PatternFill('solid', fgColor=navy)
readme.merge_cells('A45:H45')
readme['A45'] = '02_BACKLOG: të 185 task-et u rishikuan. Y = gjendja e implementimit; Z = realizuar/mbetet/burime. L ruan statusin kontraktual. 7 In Review, 106 In Progress, 35 Not Started, 2 Blocked, 35 Deferred, 0 Done.'
readme['A45'].alignment = Alignment(wrap_text=True, vertical='top')
readme.row_dimensions[45].height = 55
readme.merge_cells('A46:H46')
readme['A46'] = 'Done kërkon të gjitha acceptance criteria, provat dhe gates. Implementim lokal nuk do të thotë pranim enterprise. Shënimet/formulimet e vjetra ruhen në comments dhe backup. Shiko 14_AUDIT_2026-09-07 për provat dhe kufijtë.'
readme['A46'].alignment = Alignment(wrap_text=True, vertical='top')
readme.row_dimensions[46].height = 55

summary = wb.create_sheet('14_AUDIT_2026-09-07')
summary.sheet_properties.tabColor = '0F766E'
summary.column_dimensions['A'].width = 28
summary.column_dimensions['B'].width = 16
summary.column_dimensions['C'].width = 116
summary.merge_cells('A1:C1')
summary['A1'] = 'Rrit · Gjendja reale e 185 task-eve'
summary.merge_cells('A2:C2')
summary['A2'] = 'Audit 2026-09-07 · burimi: kodi aktual + prova lokale të ruajtura; jo deklaratë Production GA.'
summary.append([])
for col, text in enumerate(['Statusi kontraktual', 'Task-e', 'Kuptimi'], 1): summary.cell(4, col, text)
meanings = {
    'In Review': 'Implementimi kryesor i scope-it aktual ekziston; mbeten pranim/teste të plota dhe release gates. Nuk barazohet me Done.',
    'In Progress': 'Ka kod funksional ose pjesë të implementuara, por acceptance criteria kanë boshllëqe konkrete të shënuara në kolonën Z.',
    'Not Started': 'Nuk u gjet implementim/artefakt pranimi për kërkesën. Kriteri mbetet i hapur.',
    'Blocked': 'Meta App/access/credentials mungojnë; krijimi dhe verifikimi kërkon pronarin e llogarisë Meta.',
    'Deferred': 'Shtyrë nga scope-i posts-only ose zgjerim pas-GA. Detyra/ID/estimates ruhen; nuk është përfunduar.',
    'Done': 'Asnjë task i plotë origjinal nuk ka të gjitha provat dhe gates për pranim formal. Funksionaliteti real lokal është shënuar veçmas.',
}
for r, status in enumerate(meanings, 5):
    summary.cell(r, 1, status)
    summary.cell(r, 2, f'=COUNTIF(\'02_BACKLOG\'!$L$5:$L$189,A{r})')
    summary.cell(r, 3, meanings[status])
    for c in summary[r]: c.fill = PatternFill('solid', fgColor=colors[status])
summary['A11'] = 'Gjithsej'
summary['B11'] = '=SUM(B5:B10)'
summary['C11'] = '185 ID origjinale; asnjë task nuk është fshirë ose numëruar dy herë.'
notes = [
    ('Faza e përdorimit', 'Aplikacion funksional lokal me llogari/PostgreSQL dhe rrjedhë manuale të postimeve. Meta live, AI live, pagesat dhe gatishmëria enterprise nuk janë pranuar.'),
    ('Prova: unit + PostgreSQL', '33 passed në hq-optional-unit.log (2026-09-07): trial/identities/expiry, media, scope validation, token transport/signatures, schema/calendar/auth configuration.'),
    ('Prova: browser', '2 passed në hq-optional-e2e.log: Facebook dhe Instagram approve/calendar pa Meta, refresh, explicit automatic guards dhe content validation. Desktop/mobile Chromium; jo Safari/pajisje fizike.'),
    ('Prova: TypeScript / lint', 'Kontrollet e fundit të app/server types dhe lint të skedarëve të ndryshuar kaluan. Log-et: hq-optional-app-types.log, hq-optional-server-types.log, hq-optional-lint-final.log. Nuk është full production rebuild.'),
    ('Provë më e gjerë e hapur', 'live.spec.ts ka skenarë më të gjerë, por raporti i fundit i plotë nuk është green. 51 testet në docs/VALIDATION.md i përkasin demonstrimit të vjetër dhe nuk pranohen si evidencë e backend-it aktual.'),
    ('Mjedisi aktual', 'Node lokal raporton v26.5.0; .nvmrc nuk gjendet në disk. README/CI përmend Node 24. Kërkohet riprodhim i toolchain-it të synuar; konfigurimet sekrete nuk janë kopjuar në workbook.'),
    ('Konfigurime që mungojnë', 'META_APP_ID, META_APP_SECRET, META_LOGIN_CONFIG_ID, META_REDIRECT_URI, OPENAI_API_KEY, OPENAI_MODEL dhe PUBLISHING_ENABLED mungojnë. Nuk u krye publikim ose AI generation live gjatë auditit.'),
    ('Meta opsionale', 'Miratimi dhe kalendari përdorin scope content/manual. Lidhja Meta kërkohet për import dhe publikim real. Një lidhje e mëvonshme nuk e bën automatik një postim manual.'),
    ('Trial / euro / pa demo', '72 orë nga krijimi i biznesit, pa zgjatje në retry; ledger për email/browser/social identities. Browser cookie nuk është hardware fingerprint. Çmimet vijnë nga DB në EUR, por janë propozime pa checkout aktiv.'),
    ('Ndryshim arkitekture', 'Backend-i aktual është Nitro + Better Auth + PostgreSQL; plani FastAPI/Supabase dhe kërkesat RLS/outbox nuk janë realizuar siç specifikohen. AUD-2026-05 kërkon ADR; siguria nuk hiqet nga scope-i.'),
    ('Ndryshimet e workbook-ut', 'Statuset dhe shënimet e 185 task-eve; 10 task-e me scope të përshtatur dhe tekstin origjinal në comments; 19 gates me gjendje reale; vendime scope-i; counter-at S8–S10 dhe totalet u korrigjuan.'),
    ('Kopja e sigurisë', str(BACKUP.relative_to(ROOT)).replace('\\', '/')),
]
for r, (title, body) in enumerate(notes, 14):
    summary.cell(r, 1, title)
    summary.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    summary.cell(r, 2, body)
    summary.row_dimensions[r].height = 65
for row in summary:
    for c in row:
        c.font = Font(name='Aptos', size=11, color=navy, bold=c.column == 1)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = Border(bottom=line)
for r in [1, 4]:
    for c in summary[r]:
        c.fill = PatternFill('solid', fgColor=navy)
        c.font = Font(name='Aptos', size=12, bold=True, color='FFFFFF')
for r in range(5, 11): summary.row_dimensions[r].height = 47
summary.row_dimensions[1].height = 29
summary.row_dimensions[2].height = 28
summary.row_dimensions[4].height = 27
summary.freeze_panes = 'C5'
summary.sheet_view.showGridLines = False
wb.active = wb.sheetnames.index('02_BACKLOG')
wb.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)

# Supply cached values for every formula, so previewers show accurate current counts too.
# Calculations below cover this workbook's known control sheets, not arbitrary Excel code.
cache = {}
def put(sheet, coord, value): cache[(sheet, coord)] = value
kickoff = to_excel(wb['00_README']['B14'].value)
length, capacity, sprint_count = [wb['00_README'][c].value for c in ['B15', 'B16', 'B18']]
put('00_README', 'B20', kickoff + length * sprint_count - 1)
put('00_README', 'B21', kickoff + length * 9 - 1)
for r in range(5, 16):
    sprint = roadmap[f'A{r}'].value
    tasks = [ws[row] for row in rows.values() if ws.cell(row, 2).value == sprint]
    a = sum(t[10].value for t in tasks if t[8].value == 'Programmer A')
    b = sum(t[10].value for t in tasks if t[8].value == 'Programmer B')
    statuses = Counter(t[11].value for t in tasks)
    p0 = sum(t[7].value == 'P0' and t[11].value not in ['Done', 'Deferred'] for t in tasks)
    stage = ('Deferred' if statuses['Deferred'] == len(tasks) else 'Done' if statuses['Done'] == len(tasks)
             else 'In Progress' if statuses['In Progress'] + statuses['In Review'] else 'Blocked' if statuses['Blocked'] else 'Planned')
    vals = {'C': kickoff + (r-5)*length, 'D': kickoff + (r-4)*length-1, 'G': a, 'H': b, 'I': a+b,
            'J': 2*capacity, 'K': (a+b)/(2*capacity), 'L': 2*capacity-a-b, 'M': p0, 'N': len(tasks), 'O': statuses['Done'], 'P': stage}
    for col, value in vals.items(): put('01_ROADMAP', f'{col}{r}', value)
    vals = {'B': a, 'C': capacity, 'D': a/capacity, 'E': b, 'F': capacity, 'G': b/capacity,
            'H': a+b, 'I': 2*capacity, 'J': 2*capacity-a-b, 'K': p0}
    for col, value in vals.items(): put('03_WORKLOAD', f'{col}{r}', value)
for col in 'GHIJLMNO': put('01_ROADMAP', f'{col}17', sum(cache[('01_ROADMAP', f'{col}{r}')] for r in range(5, 16)))
put('01_ROADMAP', 'K17', cache[('01_ROADMAP', 'I17')] / cache[('01_ROADMAP', 'J17')])
for col in 'BEHIJK': put('03_WORKLOAD', f'{col}17', sum(cache[('03_WORKLOAD', f'{col}{r}')] for r in range(5, 16)))
for task_id, r in rows.items():
    number = int(ws.cell(r, 2).value[1:])
    put('02_BACKLOG', f'Q{r}', kickoff + number*length)
    put('02_BACKLOG', f'R{r}', kickoff + (number+1)*length-1)
for r in range(5, gates.max_row+1):
    gate = gates.cell(r, 1).value
    related = [row for row in rows.values() if gate in (ws.cell(row, 14).value or '')]
    put('04_RELEASE_GATES', f'E{r}', len(related))
    put('04_RELEASE_GATES', f'F{r}', sum(ws.cell(row, 12).value not in ['Done', 'Deferred'] for row in related))
for r in range(5, wb['05_RISKS'].max_row+1):
    score = wb['05_RISKS'].cell(r, 4).value * wb['05_RISKS'].cell(r, 5).value
    put('05_RISKS', f'F{r}', score)
    put('05_RISKS', f'G{r}', 'Critical' if score >= 15 else 'High' if score >= 10 else 'Medium' if score >= 5 else 'Low')
for r in range(5, 11): put(summary.title, f'B{r}', counts[summary.cell(r, 1).value])
put(summary.title, 'B11', sum(counts.values()))
formulas = {(s.title, c.coordinate) for s in wb for row in s for c in row if c.data_type == 'f'}
assert formulas == set(cache), (formulas - set(cache), set(cache) - formulas)
assert cache[('01_ROADMAP', 'N17')] == 185

with NamedTemporaryFile(dir=ROOT, prefix='backlog-audit-', suffix='.xlsx', delete=False) as f:
    pending = Path(f.name)
try:
    wb.save(pending)
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(pending) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(wb, 1):
        path = f'xl/worksheets/sheet{index}.xml'
        tree = ET.fromstring(files[path])
        for cell in tree.findall(f'.//{{{ns}}}c'):
            if cell.find(f'{{{ns}}}f') is None:
                continue
            value = cache[(sheet.title, cell.attrib['r'])]
            cached = cell.find(f'{{{ns}}}v')
            if cached is None: cached = ET.SubElement(cell, f'{{{ns}}}v')
            if isinstance(value, str):
                cell.set('t', 'str')
                cached.text = value
            else:
                cell.attrib.pop('t', None)
                cached.text = str(value)
        files[path] = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
    with ZipFile(pending, 'w', ZIP_DEFLATED) as archive:
        for name, content in files.items(): archive.writestr(name, content)

    verified = load_workbook(pending)
    cached = load_workbook(pending, data_only=True)
    assert verified['02_BACKLOG'].max_row == 189 and verified['02_BACKLOG'].max_column == 26
    assert [verified['02_BACKLOG'].cell(r, 1).value for r in range(5, 190)] == list(rows)
    assert all(verified['02_BACKLOG'].cell(r, 26).value for r in range(5, 190))
    assert Counter(verified['02_BACKLOG'].cell(r, 12).value for r in range(5, 190)) == counts
    assert cached['01_ROADMAP']['N17'].value == cached[summary.title]['B11'].value == 185
    for key in formulas:
        actual = cached[key[0]][key[1]].value
        if isinstance(actual, datetime):
            actual = to_excel(actual)
        assert actual == cache[key], key
    before = load_workbook(BACKUP)
    # No incidental changes outside the approved sheet/cell ranges.
    for old_sheet in before:
        after = verified[old_sheet.title]
        for row in old_sheet:
            for cell in row:
                if cell.value == after[cell.coordinate].value:
                    continue
                allowed = ((old_sheet.title == '02_BACKLOG' and cell.column in [12, 19]) or
                           (old_sheet.title, cell.coordinate) in changed_contract_cells or
                           (old_sheet.title == '01_ROADMAP' and (cell.column in [13,14,15,16] or cell.coordinate == 'A2')) or
                           (old_sheet.title == '04_RELEASE_GATES' and cell.column in [7,8,9]))
                assert allowed, (old_sheet.title, cell.coordinate, cell.value, after[cell.coordinate].value)
    for sheet, address in changed_contract_cells:
        assert verified[sheet][address].comment and str(before[sheet][address].value) in verified[sheet][address].comment.text
    assert sha256(BACKUP.read_bytes()).digest() == sha256(original_bytes).digest()
    assert TARGET.read_bytes() == original_bytes, 'Source workbook changed during audit; preserve user edits.'
    pending.replace(TARGET)
finally:
    if pending.exists(): pending.unlink()

result = {'workbook': NAME, 'date': DATE, 'task_count': 185, 'statuses': dict(counts),
          'scope_amendments': list(AMENDMENTS), 'formula_count_verified': len(formulas),
          'backup': str(BACKUP.relative_to(ROOT)), 'original_sha256': sha256(original_bytes).hexdigest(),
          'updated_sha256': sha256(TARGET.read_bytes()).hexdigest()}
(ROOT / f'docs/backlog-audit-verification-{DATE}.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(result, ensure_ascii=True, indent=2))
