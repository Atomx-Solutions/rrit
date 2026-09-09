"""Synchronize the completed S0-A03 implementation into both backlog workbooks.

The update is deliberately scoped to S0-A03 and preserves every existing formula
expression and cached formula result in the XLSX package.
"""

from copy import copy
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.datetime import to_excel


ROOT = Path(__file__).resolve().parents[1]
WORKBOOKS = (
    ROOT / "AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx",
    ROOT / "AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1 1.xlsx",
)


def formula_state(workbook, value_workbook):
    formulas = {}
    caches = {}
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                key = (sheet.title, cell.coordinate)
                formulas[key] = cell.value
                cached = value_workbook[sheet.title][cell.coordinate].value
                caches[key] = to_excel(cached) if isinstance(cached, datetime) else cached
    return formulas, caches


def restore_formula_caches(path, workbook, caches):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(path) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(workbook, 1):
        xml_path = f"xl/worksheets/sheet{index}.xml"
        tree = ET.fromstring(files[xml_path])
        for cell in tree.findall(f".//{{{namespace}}}c"):
            if cell.find(f"{{{namespace}}}f") is None:
                continue
            cached_value = caches[(sheet.title, cell.attrib["r"])]
            cached = cell.find(f"{{{namespace}}}v")
            if cached is None:
                cached = ET.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(cached_value, str):
                cell.set("t", "str")
            else:
                cell.attrib.pop("t", None)
            cached.text = "" if cached_value is None else str(cached_value)
        files[xml_path] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)


def update(path):
    original = path.read_bytes()
    backup = ROOT / "docs" / "backups" / path.name.replace(
        ".xlsx", ".before-s0-a03-2026-09-08.xlsx"
    )
    backup.parent.mkdir(parents=True, exist_ok=True)
    if not backup.exists():
        backup.write_bytes(original)

    workbook = load_workbook(path)
    values = load_workbook(path, data_only=True)
    formulas, caches = formula_state(workbook, values)
    backlog = workbook["02_BACKLOG"]
    rows = {backlog.cell(row, 1).value: row for row in range(5, backlog.max_row + 1)}
    row = rows["S0-A03"]

    updates = {
        4: "Pin Node.js 24.21.0 and npm 11.19.0 toolchain",
        5: "Apply ADR-002's Nitro/TypeScript replacement: pin runtime/package manager, enforce the npm lockfile, ESLint/type policy, deterministic Vitest settings, a versioned pre-commit hook and identical local/CI quality commands.",
        6: "Clean npm ci from lock; exact toolchain check plus lint/type/unit commands reproducible locally and in CI; Docker is out of scope for v1.",
        12: "In Review",
        13: "Clean Node 24.21.0 install: toolchain/lint/type/24 unit tests passed; integration/build represented in CI",
        15: "Dependency lock and supply-chain baseline; Supabase external database configuration.",
        16: "package manifests + package-lock + .npmrc + hooks + docs/TOOLCHAIN.md",
        19: "docs/S0-A03-REPORT-SQ.md",
        25: "Implementuar — në pranim",
        26: "ADR-002 maps the unimplemented Python backend to Nitro/TypeScript. Node.js 24.21.0/npm 11.19.0, engine enforcement, npm lockfile, shared quality commands, deterministic Vitest, pre-commit and CI parity are implemented; Docker is out of scope for v1.",
        27: "Kodi i task-ut nuk ka punë të mbetur. Ekzekuto hosted GitHub Actions; Programmer B konfirmon required check dhe regjistron RG-01 para statusit Done.",
        28: "Fresh environment build: clean npm ci; exact toolchain, lint, typecheck, unit/integration and production build; same gate in CI; reviewer/RG-01 sign-off.",
        29: "Validimi 2026-09-09: clean npm ci me Node.js 24.21.0/npm 11.19.0; toolchain/lint/typecheck kaluan; 24/24 unit tests kaluan. Hosted CI/integration/build dhe reviewer sign-off mbeten për verifikim.",
        30: ".nvmrc; .npmrc; package.json; package-lock.json; apps/web/package.json; apps/web/vitest.config.ts; .githooks/pre-commit; scripts/check-toolchain.mjs; scripts/install-git-hooks.mjs; .github/workflows/frontend.yml; docs/TOOLCHAIN.md; docs/VALIDATION.md; docs/S0-A03-REPORT-SQ.md",
        31: "Kryer në kod — pret pranim/gate",
        32: "Jo për scope-in e rreshtit",
        33: "Po — hosted CI/reviewer",
        34: "Ekzekuto hosted GitHub Actions; Programmer B konfirmon required check dhe regjistron RG-01 para statusit Done.",
        35: "docs/S0-A03-REPORT-SQ.md",
        36: "Programmer A / review: Programmer B",
    }
    for column, value in updates.items():
        cell = backlog.cell(row, column)
        cell.value = value
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    backlog.cell(row, 19).hyperlink = "docs/S0-A03-REPORT-SQ.md"
    backlog.cell(row, 35).hyperlink = "docs/S0-A03-REPORT-SQ.md"

    if "16_HANDOFF_2_DEV" in workbook.sheetnames:
        handoff = workbook["16_HANDOFF_2_DEV"]
        handoff_rows = {
            handoff.cell(candidate, 1).value: candidate
            for candidate in range(2, handoff.max_row + 1)
        }
        handoff_row = handoff_rows["S0-A03"]
        for column, value in {
            4: "In Review",
            5: "Kryer në kod — pret pranim/gate",
            8: updates[4],
            9: updates[34],
            10: "RG-01",
            11: updates[28],
        }.items():
            handoff.cell(handoff_row, column).value = value
            handoff.cell(handoff_row, column).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

    readme = workbook["00_README"]
    readme["A48"] = (
        "Code audit + S0-A03 implementation 2026-09-08: 24 task-e janë kryer në kod "
        "dhe presin pranim/gate; 90 janë pjesërisht të implementuara; 34 pa filluar; "
        "35 deferred; 2 blocked; 0 Done. Shiko 02_BACKLOG AE–AJ dhe 16_HANDOFF_2_DEV."
    )

    audit = workbook["14_AUDIT_2026-09-07"]
    audit["A32"] = "S0-A03 production-grade toolchain implementation"
    audit["B32"] = (
        "Nitro/TypeScript mapping, exact Node/npm pins, lock enforcement, deterministic tests, "
        "pre-commit and shared CI commands implemented; Docker is out of scope for v1. "
        "24 unit tests passed locally; hosted CI/reviewer RG-01 sign-off remains."
    )
    for column in (1, 2):
        audit.cell(32, column)._style = copy(audit.cell(31, column)._style)
        audit.cell(32, column).alignment = Alignment(wrap_text=True, vertical="top")
    audit.row_dimensions[32].height = 90

    current_formulas = {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook
        for sheet_row in sheet
        for cell in sheet_row
        if cell.data_type == "f"
    }
    if current_formulas != formulas:
        raise AssertionError("Formula expressions changed")

    with NamedTemporaryFile(dir=ROOT, prefix="s0-a03-", suffix=".xlsx", delete=False) as handle:
        pending = Path(handle.name)
    try:
        workbook.save(pending)
        restore_formula_caches(pending, workbook, caches)
        check = load_workbook(pending, data_only=True)
        for key, expected in caches.items():
            actual = check[key[0]][key[1]].value
            if isinstance(actual, datetime):
                actual = to_excel(actual)
            if actual != expected:
                raise AssertionError(f"Formula cache changed: {key}")
        if path.read_bytes() != original:
            raise AssertionError(f"Workbook changed concurrently: {path.name}")
        pending.replace(path)
    finally:
        if pending.exists():
            pending.unlink()

    return len(formulas), backup.relative_to(ROOT).as_posix()


for workbook_path in (candidate for candidate in WORKBOOKS if candidate.exists()):
    formula_count, backup_path = update(workbook_path)
    print(f"updated={workbook_path.name} formulas_preserved={formula_count} backup={backup_path}")
