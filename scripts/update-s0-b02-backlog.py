"""Synchronize the S0-B02 implementation evidence into the source workbook.

This script intentionally changes one ticket only. It creates a backup before
writing, preserves formula expressions and their cached values, and never
promotes the ticket to Done.
"""

from copy import copy
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
import json

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.datetime import to_excel


ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-09"
NAME = "AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx"
TARGET = ROOT / NAME
BACKUP = ROOT / "docs" / "backups" / NAME.replace(".xlsx", f".before-s0-b02-{DATE}.xlsx")
EVIDENCE = ROOT / "docs" / "backlog-s0-b02-progress-2026-09-09.json"
TASK_ID = "S0-B02"


def main() -> None:
    original = TARGET.read_bytes()
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        BACKUP.write_bytes(original)

    workbook = load_workbook(TARGET)
    values = load_workbook(TARGET, data_only=True)
    formulas = {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook
        for row in sheet
        for cell in row
        if cell.data_type == "f"
    }
    cached_values = {}
    for key in formulas:
        value = values[key[0]][key[1]].value
        if value is None:
            raise AssertionError(f"Missing cached value for {key}")
        cached_values[key] = to_excel(value) if isinstance(value, datetime) else value

    backlog = workbook["02_BACKLOG"]
    row = next(
        (index for index in range(5, backlog.max_row + 1) if backlog.cell(index, 1).value == TASK_ID),
        None,
    )
    if row is None:
        raise AssertionError(f"Could not find {TASK_ID}")
    if backlog.cell(row, 12).value not in {"In Progress", "In Review"}:
        raise AssertionError("Refusing to promote a task that is not already active")

    previous_status = backlog.cell(row, 12).value
    updates = {
        12: "In Review",
        19: "docs/S0-B02-REPORT-SQ.md",
        25: "Në review",
        26: (
            "U vendos kontrata e përbashkët `--rrit-*` për ngjyra semantike, "
            "tipografi, hapësirë, radius, hije dhe motion. Landing dhe dashboard "
            "konsumojnë tokenat e UI-së; Nuxt UI v4, Lucide dhe politika e wrapper-ave "
            "janë të dokumentuara. Light/dark tokenat dhe testet accessibility u shtuan."
        ),
        27: (
            "Mbeten GitHub Actions në PR, review i Programmer A dhe RG-14. Testi i "
            "dashboard accessibility krijon të dhëna sintetike dhe duhet të ekzekutohet "
            "vetëm në ambient të dedikuar/CI, jo kundër Supabase v1."
        ),
        28: (
            "Kriteri lokal i verifikuar: tokenat përdoren për UI standarde në landing "
            "dhe dashboard; light/dark accessibility baseline kalon për landing. "
            "Ilustrimet e landing-ut lejohen të mbajnë ngjyra/gradientë lokalë sepse "
            "janë art-direction, jo konstante të UI-së. RG-14 nuk është shënuar Passed."
        ),
        29: (
            "2026-09-09: `npm.cmd run check` kaloi (toolchain, ESLint, Nuxt typecheck, "
            "27/27 Vitest unit); `npm.cmd run build` kaloi; Playwright landing desktop "
            "kaloi 13/13 me kontroll light/dark. `git diff --check` kaloi."
        ),
        30: (
            "apps/web/app/assets/css/main.css\n"
            "apps/web/app/assets/css/landing.css\n"
            "apps/web/app/utils/design-system.ts\n"
            "apps/web/app/app.config.ts\n"
            "apps/web/tests/unit/design-system.test.ts\n"
            "apps/web/tests/e2e/accessibility-baseline.spec.ts\n"
            "docs/S0-B02-REPORT-SQ.md"
        ),
        31: "Në review — implementimi dhe validimi lokal u kryen; CI/review/RG-14 mbeten.",
    }
    for column, value in updates.items():
        cell = backlog.cell(row, column)
        cell.value = value
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    backlog.cell(row, 19).hyperlink = "docs/S0-B02-REPORT-SQ.md"
    backlog.row_dimensions[row].height = 210

    if backlog.cell(row, 12).value != "In Review":
        raise AssertionError("S0-B02 status was not updated")
    if any(
        backlog.cell(index, 12).value == "Done" and index == row
        for index in range(5, backlog.max_row + 1)
    ):
        raise AssertionError("The script must never mark S0-B02 Done")
    current_formulas = {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook
        for row_cells in sheet
        for cell in row_cells
        if cell.data_type == "f"
    }
    if formulas != current_formulas:
        raise AssertionError("Formula expressions changed")

    with NamedTemporaryFile(dir=ROOT, prefix="s0-b02-backlog-", suffix=".xlsx", delete=False) as handle:
        pending = Path(handle.name)
    try:
        workbook.save(pending)
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(pending) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        for index, sheet in enumerate(workbook, 1):
            path = f"xl/worksheets/sheet{index}.xml"
            tree = ET.fromstring(files[path])
            for cell in tree.findall(f".//{{{namespace}}}c"):
                if cell.find(f"{{{namespace}}}f") is None:
                    continue
                cached = cell.find(f"{{{namespace}}}v")
                if cached is None:
                    cached = ET.SubElement(cell, f"{{{namespace}}}v")
                value = cached_values[(sheet.title, cell.attrib["r"])]
                if isinstance(value, str):
                    cell.set("t", "str")
                else:
                    cell.attrib.pop("t", None)
                cached.text = str(value)
            files[path] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
        with ZipFile(pending, "w", ZIP_DEFLATED) as archive:
            for name, content in files.items():
                archive.writestr(name, content)

        check = load_workbook(pending)
        check_values = load_workbook(pending, data_only=True)
        if {(sheet.title, cell.coordinate): cell.value for sheet in check for row_cells in sheet for cell in row_cells if cell.data_type == "f"} != formulas:
            raise AssertionError("Formula expressions changed after writing")
        for key, expected in cached_values.items():
            actual = check_values[key[0]][key[1]].value
            if isinstance(actual, datetime):
                actual = to_excel(actual)
            if actual != expected:
                raise AssertionError(f"Formula cache changed: {key}")
        if TARGET.read_bytes() != original:
            raise AssertionError("Workbook changed concurrently")
        pending.replace(TARGET)
    finally:
        if pending.exists():
            pending.unlink()

    payload = {
        "date": DATE,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workbook": NAME,
        "task": TASK_ID,
        "previous_status": previous_status,
        "status": "In Review",
        "done_promotions": 0,
        "backup": str(BACKUP.relative_to(ROOT)).replace("\\", "/"),
        "before_sha256": sha256(original).hexdigest(),
        "after_sha256": sha256(TARGET.read_bytes()).hexdigest(),
        "formula_expressions_and_caches_preserved": len(formulas),
    }
    EVIDENCE.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
