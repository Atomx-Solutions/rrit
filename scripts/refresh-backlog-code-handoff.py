"""Reclassify code-complete tasks and build a two-developer handoff in the backlog."""

from collections import Counter
from copy import copy
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
import json
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-08"
NAME = "AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx"
TARGET = ROOT / NAME
BACKUP = ROOT / "docs" / "backups" / NAME.replace(
    ".xlsx", f".before-handoff-reclassification-{DATE}.xlsx"
)
REPORT = ROOT / "docs" / f"backlog-code-handoff-{DATE}.json"

# These tasks have their requested implementation in the repository. They still
# require the evidence/release gate documented in their own backlog row.
PROMOTE_TO_REVIEW = {
    "S0-A01": "ADR dhe lidhjet e dokumentimit janë implementuar; pret architecture review/RG-01.",
    "S0-A02": "API dhe worker-i nisen si role të ndara në container; liveness dhe container smoke kalojnë; pret RG-01.",
    "S0-A04": "PostgreSQL reset/migrimet dhe rolet migration/runtime janë implementuar e provuar; pret staging/RG-01.",
    "S0-A07": "Better Auth session guard pranon sesionin valid dhe refuzon forged/revoked/expired me kode stabile; pret RG-03.",
    "S0-A08": "Outbox-i transaksional, claim me lease/SKIP LOCKED, deduplication dhe crash recovery janë implementuar; pret RG-01.",
    "S0-B01": "Nuxt/TypeScript/runtime separation është implementuar; pret clean Node 24 build/RG-01.",
    "S0-B05": "Landing SSR/responsive/motion është implementuar; pret Lighthouse, 360px dhe content review.",
    "S0-B08": "Playwright framework, projects, trace dhe screenshots janë implementuar; pret hosted CI proof.",
    "S0-B09": "Baza e aksesueshmërisë për landing/auth/dashboard, fokusin me tastierë dhe reduced motion është implementuar e verifikuar; pret RG-14.",
    "S1-A07": "Entitlement skeleton, katalogu/limitet e tipizuara dhe trial guards janë implementuar; pret RG-08.",
    "S3-A07": "Versionet/snapshot-et immutable dhe tenant FKs janë implementuar; pret concurrency/release acceptance.",
    "S4-A04": "Facebook/Instagram publisher adapter dhe provider fixture janë implementuar; pret Meta staging proof.",
    "S4-A05": "Taxonomia e gabimeve, Retry-After, backoff me jitter, kufiri i tentimeve dhe remediation e tipizuar janë implementuar; pret Meta staging/RG-07.",
    "S4-A06": "Timeout-after-side-effect reconciliation kalon pa ridërgim; pret Meta staging/RG-07.",
    "S4-A07": "Logical key idempotent, duplicate protection dhe ledger-i append-only i tentimeve janë implementuar; pret Meta staging/RG-07.",
    "S7-B09": "Autosave/revision/conflict/navigation protection për scope-in posts-only është implementuar; pret fault E2E.",
}

EVIDENCE_UPDATES = {
    "S0-A02": (
        "Validimi 2026-09-08: image-i production me Node 24.20.0 u ngrit healthy si web dhe publisher "
        "i veçantë; /api/health/live ktheu 200 brenda container-it."
    ),
    "S0-A04": (
        "Validimi 2026-09-08: target-i migrations aplikoi migrimet 0001–0017 dhe migrimet u provuan nga zero "
        "në skemën e izoluar të integration suite. Roli runtime lexoi "
        "skemën dhe kaloi 8/8 E2E, ndërsa CREATE TABLE u refuzua me permission denied. Staging dry-run mbetet."
    ),
    "S0-A07": (
        "Validimi 2026-09-08: E2E mbi container-in production provoi sesion valid, cookie të falsifikuar, "
        "logout/revokim dhe expiresAt të kaluar; përgjigjja e mbrojtur ktheu AUTH_REQUIRED. Origin-i i huaj "
        "u refuzua me ORIGIN_REJECTED."
    ),
    "S0-A08": (
        "Validimi 2026-09-08: migrimi 0014 krijoi JobEnvelope/outbox receipts; testet provuan commit/rollback "
        "atomik, dedupe, dy consumer-a paralelë me vetëm një efekt dhe rikuperim të claim-it stale. Readiness "
        "production kërkon migrimin minimal aktual 0017."
    ),
    "S0-A09": (
        "Validimi 2026-09-08: workflow-i kaloi actionlint dhe tani ndërton/nis image-in e ngurtësuar për E2E; "
        "required-check/merge-blocking në GitHub mbetet provë e jashtme."
    ),
    "S0-B01": (
        "Validimi 2026-09-08: clean npm ci, nuxt typecheck dhe production build kaluan brenda image-it "
        "Node 24.20.0 të fiksuar me digest."
    ),
    "S0-B07": (
        "Validimi 2026-09-08: CI ndërton container-in production dhe E2E ekzekutohet mbi të; actionlint kaloi. "
        "Hosted CI dhe branch protection mbeten gate të jashtme."
    ),
    "S0-B08": (
        "Validimi 2026-09-08: 14/14 E2E kaluan në Chromium desktop/mobile kundër container-it production; "
        "workflow-i ruan trace, screenshot, report dhe production log."
    ),
    "S0-B09": (
        "Validimi 2026-09-08: Axe WCAG 2 A/AA nuk gjeti çështje serious/critical në landing, login, "
        "signup ose dashboard; fokusi i tastierës ishte i dukshëm, reduced-motion kufizoi animacionet dhe "
        "layout-i nuk pati overflow në Chromium desktop/mobile. 4/4 prova të dedikuara dhe 14/14 E2E kaluan "
        "kundër image-it production."
    ),
    "S4-A05": (
        "Validimi 2026-09-08: migrimi 0015 ruan retry_count/max_attempts për çdo job; fault fixtures mbulojnë "
        "429/Retry-After, 5xx, network, auth, permission dhe media. Integration test provoi backoff, jitter, "
        "kufirin tre-tentimësh dhe terminal code; shkrimi me rezultat të paqartë vazhdon në reconciliation pa "
        "ridërgim. API kthen disposition/remediation të tipizuar dhe UI shfaq retry/blocked/terminal state."
    ),
    "S4-A07": (
        "Validimi 2026-09-08: migrimet 0016–0017 shtuan ledger-in append-only, attempt sequence, request hash "
        "dhe logical key unik për biznes. Dy queue calls konkurruese kthyen të njëjtin job; dy worker-a prodhuan "
        "vetëm një provider POST; retry-t ruajtën të njëjtin hash dhe histori të renditur. UPDATE/DELETE u refuzuan "
        "nga trigger-i dhe roli runtime ka vetëm SELECT/INSERT mbi ledger-in."
    ),
    "S8-B01": (
        "Validimi 2026-09-08: 14/14 critical E2E kaluan kundër artefaktit production në Chromium desktop/mobile; "
        "Firefox/WebKit dhe pajisjet e tjera të mbështetura mbeten."
    ),
}

PATH_UPDATES = {
    "S0-A02": "Dockerfile; compose.yaml; apps/web/server/plugins/publisher.ts; apps/web/server/api/health/live.get.ts; docs/DEPLOYMENT.md",
    "S0-A04": "Dockerfile; compose.yaml; infra/postgres/provision-runtime.sh; scripts/setup-local.mjs; apps/web/scripts/migrate.mjs; apps/web/database/migrations; docs/DEPLOYMENT.md",
    "S0-A07": "apps/web/server/utils/auth.ts; apps/web/server/utils/access.ts; apps/web/server/utils/auth-configuration.ts; apps/web/tests/e2e/auth-session.spec.ts; apps/web/tests/e2e/local-registration.spec.ts",
    "S0-A08": "apps/web/database/migrations/0014_durable_outbox.sql; apps/web/server/utils/outbox.ts; apps/web/tests/integration/server.test.ts; apps/web/server/api/health/ready.get.ts",
    "S0-A09": ".github/workflows/frontend.yml; Dockerfile; docs/VALIDATION.md",
    "S0-B01": "Dockerfile; .nvmrc; package-lock.json; apps/web/package.json; docs/VALIDATION.md",
    "S0-B07": ".github/workflows/frontend.yml; Dockerfile; apps/web/scripts/bundle-budget.mjs",
    "S0-B08": ".github/workflows/frontend.yml; apps/web/playwright.config.ts; apps/web/tests/e2e; docs/VALIDATION.md",
    "S0-B09": "apps/web/tests/e2e/accessibility-baseline.spec.ts; apps/web/app/assets/css/main.css; apps/web/app/components/AuthSurface.vue; docs/VALIDATION.md",
    "S4-A05": "apps/web/database/migrations/0015_publication_retry_policy.sql; apps/web/server/utils/meta.ts; apps/web/server/utils/publishing.ts; apps/web/server/api/posts/[id]/publication.get.ts; apps/web/app/components/ContentEditor.vue; apps/web/tests/integration/server.test.ts",
    "S4-A07": "apps/web/database/migrations/0016_publication_attempt_ledger.sql; apps/web/database/migrations/0017_publication_logical_keys.sql; infra/postgres/provision-runtime.sh; apps/web/server/utils/publishing.ts; apps/web/server/api/posts/[id]/publish.post.ts; apps/web/server/api/posts/[id]/publication.get.ts; apps/web/tests/integration/server.test.ts",
    "S8-B01": "apps/web/playwright.config.ts; apps/web/tests/e2e; docs/VALIDATION.md",
}

EXTERNAL_TERMS = (
    "meta app",
    "meta staging",
    "provider",
    "smtp",
    "hosted ci",
    "uat",
    "pajisje fizike",
    "legal",
    "ligjor",
    "komercial",
    "review",
    "release gate",
    "rg-",
    "staging",
    "production",
)

SOURCE_SUFFIXES = {
    ".ts",
    ".vue",
    ".mjs",
    ".js",
    ".sql",
    ".css",
    ".html",
    ".json",
    ".md",
    ".yaml",
    ".yml",
    ".sh",
}
EXCLUDED_PARTS = {
    "node_modules",
    ".nuxt",
    ".output",
    "test-results",
    "playwright-report",
    "__pycache__",
    "backups",
}


def source_inventory():
    roots = [ROOT / "apps" / "web", ROOT / ".github", ROOT / "docs", ROOT / "scripts", ROOT / "infra"]
    items = []
    candidates = [
        ROOT / "Dockerfile",
        ROOT / "compose.yaml",
        ROOT / ".dockerignore",
        ROOT / ".env.example",
        ROOT / ".nvmrc",
        ROOT / "package.json",
    ]
    for base in roots:
        if not base.exists():
            continue
        candidates.extend(base.rglob("*"))
    for path in candidates:
            relative = path.relative_to(ROOT)
            if (
                not path.is_file()
                or any(part in EXCLUDED_PARTS for part in relative.parts)
                or (path.suffix.lower() not in SOURCE_SUFFIXES and path.name not in {"Dockerfile", ".dockerignore", ".env.example", ".nvmrc"})
                or path.name.endswith(".log")
                or path.name.startswith("backlog-") and path.suffix == ".json"
                or path.name.startswith("deep-audit-")
            ):
                continue
            raw = path.read_bytes()
            text = raw.decode("utf-8", errors="replace")
            items.append(
                {
                    "path": relative.as_posix(),
                    "bytes": len(raw),
                    "lines": len(text.splitlines()),
                    "sha256": sha256(raw).hexdigest(),
                }
            )
    return sorted(items, key=lambda item: item["path"])


def first_action(value):
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    text = re.sub(r"^\d+\.\s*", "", text)
    text = re.split(r"\s+/\s+\d+\.\s+|\s+DA-\d+:", text, maxsplit=1)[0]
    return text[:700] or "Lexo acceptance criteria dhe përcakto implementimin/provën e munguar."


original = TARGET.read_bytes()
BACKUP.parent.mkdir(parents=True, exist_ok=True)
if not BACKUP.exists():
    BACKUP.write_bytes(original)

wb = load_workbook(TARGET)
value_wb = load_workbook(TARGET, data_only=True)
formulas = {
    (sheet.title, cell.coordinate): cell.value
    for sheet in wb
    for row in sheet
    for cell in row
    if cell.data_type == "f"
}
formula_cache = {}
for key in formulas:
    value = value_wb[key[0]][key[1]].value
    if value is None:
        raise AssertionError(f"Missing formula cache for {key}")
    formula_cache[key] = to_excel(value) if isinstance(value, datetime) else value

ws = wb["02_BACKLOG"]
rows = {ws.cell(row, 1).value: row for row in range(5, ws.max_row + 1)}
if not PROMOTE_TO_REVIEW.keys() <= rows.keys():
    raise AssertionError("Promotion task ID mismatch")

# The product decision replaced the original Supabase/FastAPI bootstrap with a
# PostgreSQL/Nitro modular monolith. Keep the task's intent and acceptance proof
# aligned with the architecture that is actually shipped.
postgres_row = rows["S0-A04"]
for column, value in {
    4: "Bootstrap PostgreSQL environments and migrations",
    5: "Create local/staging/prod conventions, versioned SQL migrations, separate migration/runtime database roles and a repeatable empty-database strategy.",
    6: "Local reset succeeds from zero; migration/runtime roles are separated; staging dry-run documented; no production credentials committed.",
    13: "PostgreSQL empty-database reset + migration/runtime role smoke",
    16: "Docker/ migrations + role provisioning + deploy runbook",
    25: "Implementuar — në pranim",
    26: "PostgreSQL në Docker, migrime SQL me checksum/transaksion/advisory lock, target i veçantë migrimi dhe rol runtime pa DDL.",
    27: "Ekzekuto migration dry-run në staging me role të ndara dhe regjistro RG-01; kjo kërkon mjedisin e hostuar.",
    28: "Prova e kërkuar: empty-database migration smoke, runtime DDL-deny, staging dry-run dhe RG-01.",
}.items():
    ws.cell(postgres_row, column).value = value

auth_row = rows["S0-A07"]
for column, value in {
    3: "Authentication",
    4: "Verify Better Auth sessions in Nitro",
    5: "Validate opaque database sessions through Better Auth on every protected Nitro endpoint; enforce expiry, revocation, forged-cookie rejection and exact-origin checks for mutations.",
    6: "Valid session accepted; expired, revoked and forged session rejected with AUTH_REQUIRED; untrusted mutation origin rejected with ORIGIN_REJECTED.",
    13: "Production-container auth session E2E",
    16: "Central session/origin guard + auth acceptance suite",
    25: "Implementuar — në pranim",
    26: "Better Auth me sessione opaque në PostgreSQL, cookie cache të çaktivizuar, skadim 7-ditor, revokim pas logout/reset dhe guard i përbashkët për API-të.",
    27: "Ekzekuto suite-n në hosted CI/public staging dhe regjistro security review/RG-03.",
    28: "Prova e kërkuar: valid/forged/revoked/expired session E2E, origin rejection dhe RG-03.",
}.items():
    ws.cell(auth_row, column).value = value

outbox_row = rows["S0-A08"]
for column, value in {
    25: "Implementuar — në pranim",
    26: "Outbox PostgreSQL me envelope/version/dedupe, claim FOR UPDATE SKIP LOCKED, lease recovery, receipt të konsumit dhe transaksion të njëjtë për efektet DB.",
    27: "Ekzekuto hosted CI/container failure rehearsal dhe regjistro RG-01; adapterët me side-effect të jashtëm vazhdojnë të përdorin idempotency/reconciliation specifik.",
    28: "Prova e kërkuar: queue duplicate/crash/concurrency tests, container smoke dhe RG-01.",
}.items():
    ws.cell(outbox_row, column).value = value

promoted = []
for task_id in PROMOTE_TO_REVIEW:
    row = rows[task_id]
    status = ws.cell(row, 12).value
    if status not in {"In Progress", "In Review"}:
        raise AssertionError(f"Unexpected status for {task_id}: {status}")
    if status == "In Progress":
        ws.cell(row, 12).value = "In Review"
        promoted.append(task_id)
    ws.cell(row, 25).value = "Implementuar — në pranim"
    classification = "Klasifikimi 2026-09-08: " + PROMOTE_TO_REVIEW[task_id]
    implemented = str(ws.cell(row, 26).value or "").rstrip()
    if classification not in implemented:
        ws.cell(row, 26).value = implemented + "\n\n" + classification

for task_id, evidence in EVIDENCE_UPDATES.items():
    row = rows[task_id]
    current = str(ws.cell(row, 29).value or "").rstrip()
    if evidence not in current:
        ws.cell(row, 29).value = current + "\n\n" + evidence
    paths = str(ws.cell(row, 30).value or "").rstrip()
    path_update = PATH_UPDATES[task_id]
    if path_update not in paths:
        ws.cell(row, 30).value = paths + "\n\n" + path_update

headers = {
    31: "Gjendja reale e kodit — 2026-09-08",
    32: "Kërkon zhvillim tjetër?",
    33: "Kërkon provë / gate?",
    34: "Hapi i ardhshëm konkret",
    35: "Evidenca e auditit të ri",
    36: "Pronari / reviewer për handoff",
}
for column, label in headers.items():
    cell = ws.cell(4, column)
    cell.value = label
    cell._style = copy(ws.cell(4, 30)._style)
    cell.alignment = Alignment(wrap_text=True, vertical="center")

inventory = source_inventory()
inventory_hash = sha256(
    "\n".join(f"{item['path']}:{item['sha256']}" for item in inventory).encode()
).hexdigest()
task_records = []
for row in range(5, ws.max_row + 1):
    task_id = ws.cell(row, 1).value
    if not task_id:
        continue
    status = ws.cell(row, 12).value
    if status == "In Review":
        code_state = "Kryer në kod — pret pranim/gate"
        development = "Jo për scope-in e rreshtit"
        gate = "Po"
    elif status == "In Progress":
        code_state = "Pjesërisht — kërkon zhvillim"
        development = "Po"
        gate = "Po, pas implementimit"
    elif status == "Not Started":
        code_state = "Pa filluar"
        development = "Po"
        gate = "Po, pas implementimit"
    elif status == "Deferred":
        code_state = "Shtyrë nga scope-i aktual"
        development = "Jo në scope-in posts-only"
        gate = "Jo tani"
    elif status == "Blocked":
        code_state = "Bllokuar nga varësi e jashtme"
        development = "Po, pas zhbllokimit"
        gate = "Po — pronari/provider-i"
    elif status == "Done":
        code_state = "Kryer dhe pranuar"
        development = "Jo"
        gate = "Kaluar"
    else:
        raise AssertionError(f"Unknown status {status} for {task_id}")
    remaining = str(ws.cell(row, 27).value or "")
    external = any(term in remaining.lower() for term in EXTERNAL_TERMS)
    if external and status == "In Progress":
        gate = "Po — përzier me punë të jashtme"
    ws.cell(row, 31).value = code_state
    ws.cell(row, 32).value = development
    ws.cell(row, 33).value = gate
    ws.cell(row, 34).value = first_action(remaining)
    ws.cell(row, 35).value = f"docs/backlog-code-handoff-{DATE}.json#{task_id}"
    ws.cell(row, 35).hyperlink = f"docs/backlog-code-handoff-{DATE}.json"
    ws.cell(row, 36).value = f"{ws.cell(row, 9).value} / review: {ws.cell(row, 10).value}"
    for column in range(31, 37):
        ws.cell(row, column).alignment = Alignment(wrap_text=True, vertical="top")
    task_records.append(
        {
            "task_id": task_id,
            "sprint": ws.cell(row, 2).value,
            "title": ws.cell(row, 4).value,
            "formal_status": status,
            "code_state": code_state,
            "owner": ws.cell(row, 9).value,
            "reviewer": ws.cell(row, 10).value,
            "priority": ws.cell(row, 8).value,
            "release_gate": ws.cell(row, 14).value,
            "next_action": ws.cell(row, 34).value,
            "required_evidence": ws.cell(row, 28).value,
            "observed_evidence": ws.cell(row, 29).value,
            "paths": ws.cell(row, 30).value,
        }
    )

for column, width in {31: 34, 32: 24, 33: 28, 34: 74, 35: 42, 36: 34}.items():
    ws.column_dimensions[get_column_letter(column)].width = width
ws.auto_filter.ref = f"A4:AJ{ws.max_row}"

sheet_name = "16_HANDOFF_2_DEV"
if sheet_name in wb.sheetnames:
    del wb[sheet_name]
handoff = wb.create_sheet(sheet_name)
handoff_headers = [
    "Task ID",
    "Sprint",
    "Priority",
    "Status formal",
    "Gjendja e kodit",
    "Owner",
    "Reviewer",
    "Task",
    "Hapi i ardhshëm",
    "Release gate",
    "Prova e kërkuar",
]
for column, label in enumerate(handoff_headers, 1):
    cell = handoff.cell(1, column, label)
    cell._style = copy(ws.cell(4, min(column, 30))._style)
    cell.alignment = Alignment(wrap_text=True, vertical="center")

status_order = {"In Review": 0, "In Progress": 1, "Blocked": 2, "Not Started": 3}
actionable = [record for record in task_records if record["formal_status"] != "Deferred"]
actionable.sort(
    key=lambda record: (
        status_order.get(record["formal_status"], 9),
        str(record["priority"]),
        record["task_id"],
    )
)
for row, record in enumerate(actionable, 2):
    values = [
        record["task_id"],
        record["sprint"],
        record["priority"],
        record["formal_status"],
        record["code_state"],
        record["owner"],
        record["reviewer"],
        record["title"],
        record["next_action"],
        record["release_gate"],
        record["required_evidence"],
    ]
    for column, value in enumerate(values, 1):
        cell = handoff.cell(row, column, value)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    handoff.row_dimensions[row].height = 72
for column, width in enumerate([14, 9, 10, 16, 32, 18, 18, 42, 78, 18, 72], 1):
    handoff.column_dimensions[get_column_letter(column)].width = width
handoff.freeze_panes = "A2"
handoff.auto_filter.ref = f"A1:K{handoff.max_row}"

audit = wb["14_AUDIT_2026-09-07"]
audit["A31"] = "Riklasifikimi code-complete dhe handoff për dy programerë"
audit["B31"] = (
    f"{len(PROMOTE_TO_REVIEW)} task-e u riklasifikuan In Progress→In Review pas krahasimit me kodin aktual. "
    f"U indeksuan {len(inventory)} skedarë first-party; 02_BACKLOG ka kolonat AE–AJ dhe "
    "16_HANDOFF_2_DEV ka listën e punës. Done mbetet vetëm për task-e me gate të kaluar."
)
for column in (1, 2):
    audit.cell(31, column)._style = copy(audit.cell(30, column)._style)
    audit.cell(31, column).alignment = Alignment(wrap_text=True, vertical="top")
audit.row_dimensions[31].height = 90

formal_counts = Counter(record["formal_status"] for record in task_records)
code_counts = Counter(record["code_state"] for record in task_records)
expected = {
    "In Review": 23,
    "In Progress": 90,
    "Not Started": 35,
    "Deferred": 35,
    "Blocked": 2,
}
if dict(formal_counts) != expected:
    raise AssertionError(f"Unexpected reclassification counts: {formal_counts}")
if len(task_records) != 185 or len({record['task_id'] for record in task_records}) != 185:
    raise AssertionError("Task inventory is incomplete or duplicated")
if formulas != {
    (sheet.title, cell.coordinate): cell.value
    for sheet in wb
    for row in sheet
    for cell in row
    if cell.data_type == "f"
}:
    raise AssertionError("Formula expressions changed")

wb["00_README"]["A48"] = (
    "Code audit 2026-09-08: 23 task-e janë kryer në kod dhe presin pranim/gate; "
    "90 janë pjesërisht të implementuara; 35 pa filluar; 35 deferred; 2 blocked; 0 Done. "
    "Shiko 02_BACKLOG AE–AJ dhe 16_HANDOFF_2_DEV."
)

with NamedTemporaryFile(dir=ROOT, prefix="backlog-handoff-", suffix=".xlsx", delete=False) as handle:
    pending = Path(handle.name)
try:
    wb.save(pending)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(pending) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(wb, 1):
        path = f"xl/worksheets/sheet{index}.xml"
        tree = ET.fromstring(files[path])
        for cell in tree.findall(f".//{{{namespace}}}c"):
            if cell.find(f"{{{namespace}}}f") is None:
                continue
            cached_value = formula_cache[(sheet.title, cell.attrib["r"])]
            cached = cell.find(f"{{{namespace}}}v")
            if cached is None:
                cached = ET.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(cached_value, str):
                cell.set("t", "str")
            else:
                cell.attrib.pop("t", None)
            cached.text = str(cached_value)
        files[path] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with ZipFile(pending, "w", ZIP_DEFLATED) as archive:
        for path, content in files.items():
            archive.writestr(path, content)
    check_values = load_workbook(pending, data_only=True)
    if len(formulas) != 706:
        raise AssertionError("Unexpected formula count")
    for key, expected_value in formula_cache.items():
        actual = check_values[key[0]][key[1]].value
        if isinstance(actual, datetime):
            actual = to_excel(actual)
        if actual != expected_value:
            raise AssertionError(f"Formula cache changed: {key}")
    if TARGET.read_bytes() != original:
        raise AssertionError("Workbook changed concurrently")
    pending.replace(TARGET)
finally:
    if pending.exists():
        pending.unlink()

payload = {
    "date": DATE,
    "generated_at": datetime.now(timezone.utc).isoformat(),
    "workbook": NAME,
    "audit_scope": "all 185 backlog rows against the current first-party source inventory",
    "source_files_read": len(inventory),
    "source_lines_read": sum(item["lines"] for item in inventory),
    "source_inventory_hash": inventory_hash,
    "source_inventory": inventory,
    "reclassified_as_code_complete": sorted(PROMOTE_TO_REVIEW),
    "newly_promoted_on_this_run": promoted,
    "formal_counts": dict(formal_counts),
    "code_counts": dict(code_counts),
    "done_policy": "Done only after row acceptance criteria, required test evidence and release gate pass.",
    "tasks": task_records,
    "backup": BACKUP.relative_to(ROOT).as_posix(),
    "before_sha256": sha256(original).hexdigest(),
    "after_sha256": sha256(TARGET.read_bytes()).hexdigest(),
    "formula_expressions_and_caches_preserved": len(formulas),
}
REPORT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(
    json.dumps(
        {
            "reclassified_as_code_complete": sorted(PROMOTE_TO_REVIEW),
            "newly_promoted_on_this_run": promoted,
            "formal_counts": dict(formal_counts),
            "code_counts": dict(code_counts),
            "source_files_read": len(inventory),
            "source_lines_read": payload["source_lines_read"],
            "formula_expressions_and_caches_preserved": len(formulas),
            "backup": payload["backup"],
        },
        ensure_ascii=False,
        indent=2,
    )
)
