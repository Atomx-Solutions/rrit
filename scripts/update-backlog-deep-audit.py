"""Apply the second, source-wide audit without changing task contracts or estimates.

Run from any directory with Python + openpyxl. The first audit must already exist.
Preserves all 706 formula expressions AND their checked cached values. Writes only
the workbook and audit documents, never application records or configuration.
"""
from collections import Counter
from copy import copy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.views import Selection

ROOT = Path(__file__).resolve().parents[1]
DATE = '2026-09-07'
NAME = 'AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx'
TARGET = ROOT / NAME
BACKUP = ROOT / 'docs/backups' / NAME.replace('.xlsx', f'.before-deep-audit-{DATE}.xlsx')
def read(name):
    return json.loads((ROOT / 'docs' / name).read_text(encoding='utf-8'))
def write(name, value):
    (ROOT / 'docs' / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

first = read(f'backlog-audit-{DATE}.json')
remaining = read(f'backlog-deep-remaining-{DATE}.json')
findings = read(f'backlog-deep-findings-{DATE}.json')
inventory = read('deep-audit-file-inventory.json')
coverage = read('deep-audit-read-coverage.json')
assert len(inventory) == len(coverage) == 133
assert {x['path'] for x in inventory} == set(coverage)
source_changes = []
for item in inventory:
    path = item['path']
    current = (ROOT / path).read_bytes()
    digest = sha256(current).hexdigest()
    if digest != coverage[path]['sha256']:
        # The only reviewed source change in this audit is the pagination locator.
        assert path == 'apps/web/tests/e2e/live.spec.ts', f'Unexpected source change: {path}'
        source_changes.append({'path': path, 'before': coverage[path]['sha256'], 'after': digest,
                               'reason': 'DA-17: select the visible Next Page control; retain result assertions.'})
        coverage[path]['previous_sha256'] = coverage[path]['sha256']
        coverage[path]['sha256'] = digest
    item['lines'] = coverage[path]['lines'] = len(current.decode('utf-8-sig').splitlines())
    item['bytes'] = len(current)
for finding in findings:
    assert set(finding['tasks']) <= set(first)
    for path in finding['paths']:
        assert (ROOT / re.sub(r':\d+$', '', path)).exists(), path

original = TARGET.read_bytes()
if BACKUP.exists():
    assert BACKUP.read_bytes() == original, 'Workbook differs from saved first audit; do not overwrite newer work.'
else:
    BACKUP.write_bytes(original)
wb = load_workbook(TARGET)
old_values = load_workbook(TARGET, data_only=True)
ws = wb['02_BACKLOG']
assert ws.max_column == 26 and ws.max_row == 189
rows = {ws.cell(r, 1).value: r for r in range(5, 190)}
assert set(rows) == set(first) and len(rows) == 185
assert set(remaining) <= set(rows)
original_formulas = {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == 'f'}
cache = {}
for key in original_formulas:
    value = old_values[key[0]][key[1]].value
    assert value is not None, f'Missing baseline formula cache: {key}'
    cache[key] = to_excel(value) if isinstance(value, datetime) else value

evidence = {}
def proof(ids, detail):
    for task in ids.split():
        evidence.setdefault(task, []).append(detail)
proof('S1-A07 S5-A06 S5-B06 S7-A04 S7-A05',
      'hq-deep-audit-unit.log: trial 72h/idempotency, concurrent device claim, normalized email, reused Meta identity, expiry dhe active entitlement. Nuk provon checkout, subscription lifecycle ose hardware identity.')
proof('S0-A04 S2-A02 S2-A03 S2-A05 S2-A04 S1-A09 S5-A08',
      'hq-deep-audit-unit.log: schema testuese PostgreSQL nga migrimet, JPEG decode/EXIF, ownership checks, token encryption/transport dhe signed media. Nuk është RLS/ASVS/Meta OAuth acceptance e plotë.')
proof('S4-A01 S4-A02 S4-A07 S4-B01 S4-B02 S4-B03 S7-B03 S7-A03',
      'hq-deep-audit-unit.log + hq-deep-audit-registration-meta-e2e.log: transition guards, content/publication scope, manual calendar pa Meta, automatic explicit, DST utility. Nuk provon worker publish real ose Approver role matrix.')
proof('S0-B04 S1-B01 S1-B02 S8-B12 S8-B13',
      'hq-deep-audit-registration-meta-e2e.log: regjistrim pa email lokal dhe onboarding; Facebook/Instagram approval/calendar pa Meta në Chromium desktop/mobile. Nuk është UAT ose email delivery publik.')
proof('S0-B03 S0-B09 S1-B08 S2-B09 S3-B01 S3-B02 S4-B08 S5-B08 S6-B10 S8-B01 S8-B02',
      'hq-deep-audit-live-e2e-rerun.log: 4/4. Business/media/validation/revision conflict dhe pagination 13 postime; axe dhe overflow në overview/content/gallery/calendar/social-posts/billing/settings-security, 1440px/390px Chromium. Nuk është Firefox/WebKit/real-device/fault/UAT matrix.')
proof('S0-B05 S5-B07',
      'hq-deep-audit-live-e2e-rerun.log: faqet publike dhe rrugët drejt regjistrimit kalojnë në desktop/mobile. Lighthouse/Web Vitals në build production nuk u matën.')
proof('S0-B01 S0-B02 S0-B06 S0-B07 S0-A09',
      'hq-deep-audit-lint.log, hq-deep-audit-app-types.log, hq-deep-audit-server-types.log: exit 0; lint i plotë dhe TypeScript app/server. Nuk u krye clean install/build prodhimi ose hosted CI.')
proof('S0-B08',
      'DA-17 u riprodhua: run i parë 2 passed/2 failed; u korrigjua vetëm selektori Next Page, rerun 4/4. Registration/optional-meta: 4/4. Cleanup dhe hosted CI mbeten të hapura.')
proof('S5-B03',
      'hq-deep-audit-unit.log: formatim EUR; live E2E hap billing. Nuk ka contract snapshot test që pajton çdo çmim/feature/limit me serverin.')
proof('S4-A04 S4-A09 S8-A11',
      'Suite aktuale u lexua; nuk ekzekuton runPublishingBatch me matricën e plotë të efekteve Meta. Asnjë postim real ose test live provider nuk u dërgua.')

deep = {}
for task_id, r in rows.items():
    entry = first[task_id]
    assert entry['status'] == ws.cell(r, 12).value
    linked = [f for f in findings if task_id in f['tasks']]
    missing = [remaining.get(task_id, entry['remaining'])]
    missing += [f"{f['id']}: {f['remaining']}" for f in linked]
    close = [f"Prova e kërkuar nga backlog-u: {ws.cell(r, 13).value}",
             f"Kriteri për t'u verifikuar (nuk deklarohet i kaluar): {ws.cell(r, 6).value}"]
    close += [f"{f['id']}: {f['verification']}" for f in linked]
    if ws.cell(r, 14).value:
        close.append(f"Pranimi kërkon edhe gate-et: {ws.cell(r, 14).value}. Asnjë gate nuk është shënuar Passed nga ky audit.")
    actual = evidence.get(task_id, [
        'U rishikuan kodi/artefaktet dhe u kërkua implementimi i kësaj detyre. Nuk ka provë të ekzekutuar specifike që mbyll të gjithë kriterin; kontrollet globale nuk e zëvendësojnë atë.'
    ])
    if entry['status'] == 'Deferred':
        actual = ['Shtyrë sipas scope-it; nuk është realizuar. Kriteret dhe hapat e këtij rreshti vlejnë kur moduli të rikthehet në scope.']
    if entry['status'] == 'Blocked':
        actual = ['Varësi e jashtme Meta App/access e pazgjidhur. Nuk ka provë live me biznes jo-developer; kod OAuth i pranishëm nuk e mbyll App Review.']
    paths = list(dict.fromkeys(entry['paths'] + [p for f in linked for p in f['paths']]))
    deep[task_id] = {
        'title': ws.cell(r, 4).value, 'status': entry['status'], 'implemented': entry['done'],
        'remaining': missing, 'closure_evidence_required': close,
        'evidence_observed': actual, 'findings': [f['id'] for f in linked], 'paths': paths,
        'audit_method': 'first-party source review + scoped local tests; not a production sign-off',
    }
    ws.cell(r, 19).value = f'docs/backlog-deep-audit-{DATE}.json#{task_id}'
    ws.cell(r, 19).hyperlink = f'docs/backlog-deep-audit-{DATE}.json'
    ws.cell(r, 19).comment = Comment('Burim i rishikimit, jo deklaratë se Test Evidence M është kaluar.', 'Codex')
    ws.cell(r, 26).value = entry['done']
    ws.cell(r, 27).value = '\n'.join(f'{i}. {text}' for i, text in enumerate(missing, 1))
    ws.cell(r, 28).value = '\n\n'.join(close)
    ws.cell(r, 29).value = '\n\n'.join(actual)
    ws.cell(r, 30).value = '\n'.join(paths + [f"{f['id']}: {f['finding']}" for f in linked]) or 'Nuk u gjet implementim specifik; shiko kriterin dhe hapat në AA/AB.'
    for col in range(26, 31):
        c = ws.cell(r, col)
        c._style = copy(ws.cell(r, 26)._style)
        c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[r].height = min(360, max(155, 75 + len(missing) * 42))

headers = {25: 'Vlerësimi i zbatimit', 26: 'Çfarë është kryer realisht',
           27: 'Çfarë mbetet për këtë task', 28: 'Prova / kushti për mbyllje',
           29: 'Provat e kryera dhe kufijtë', 30: 'Kodi / gjetjet e auditit'}
for col, label in headers.items():
    ws.cell(4, col).value = label
    ws.cell(4, col)._style = copy(ws['Y4']._style)
    ws.cell(4, col).alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[get_column_letter(col)].width = {25: 26, 26: 64, 27: 78, 28: 78, 29: 70, 30: 76}[col]
ws['AA4'].comment = Comment('Hapa të mbetur për çdo task; DA-xx lidhen me sheet 15_GJETJET. Deferred nuk do të thotë Done.', 'Codex')
ws['AB4'].comment = Comment('Testet/kriteret e pritshme. Rezultatet e ekzekutuara janë vetëm në AC.', 'Codex')
ws.auto_filter.ref = 'A4:AD189'
ws.freeze_panes = 'B5'
ws.sheet_view.pane.topLeftCell = 'Y5'
ws.sheet_view.selection = [Selection(pane='topRight', activeCell='Y1', sqref='Y1'),
                           Selection(pane='bottomLeft', activeCell='A5', sqref='A5'),
                           Selection(pane='bottomRight', activeCell='Y5', sqref='Y5')]
ws.sheet_view.zoomScale = 75

summary = wb['14_AUDIT_2026-09-07']
summary['A2'] = 'Rishikim i thelluar 2026-09-07 · 133 skedarë / 185 task-e · prova lokale; jo pranim Production GA.'
summary['C6'] = 'Ka implementim të pjesshëm; Z tregon realizimin, AA hapat e mbetur, AB provën e kërkuar, AC rezultatin aktual dhe AD kodin/gjetjet.'
updates = {
    15: ('Prova: unit + PostgreSQL', '33/33 kaluan: hq-deep-audit-unit.log. Schema testuese nga migrimet, trial/identities/expiry, media, validation, crypto/transport, workspace/DST/auth config.'),
    16: ('Prova: browser', '8/8 në dy suite: live rerun 4/4 dhe registration/optional-meta 4/4. Chromium 1440px/390px, përfshirë axe në 7 faqe dashboard-i. Log-et: hq-deep-audit-live-e2e-rerun.log dhe hq-deep-audit-registration-meta-e2e.log.'),
    17: ('Prova: TypeScript / lint', 'Full eslint . dhe TypeScript app/server kaluan: hq-deep-audit-lint.log, hq-deep-audit-app-types.log, hq-deep-audit-server-types.log. Pas DA-17 kaloi edhe eslint i live.spec.ts.'),
    18: ('Dështimi i regjistruar dhe korrigjimi', 'live.spec.ts fillimisht 2 passed/2 failed: kërkonte numrin 2 kur pamja kompakte shfaq vetëm numrin aktual. DA-17 korrigjon selektorin te Next Page pa hequr assertions; rerun 4/4. Raporti i vjetër demo në VALIDATION.md nuk zëvendëson këto rezultate.'),
    24: ('Leximi dhe ndryshimet e auditit', f"133 skedarë first-party, {sum(x['lines'] for x in inventory)} rreshta, përfshirë 7 migrimet, server/app/shared/test/config. 185 task-e me realizim/mbetje/prova veçmas; 17 gjetje DA-01..DA-17. Ndryshimi i vetëm në kod është selektori i testit DA-17."),
    25: ('Backup para rishikimit të dytë', str(BACKUP.relative_to(ROOT)).replace('\\', '/')),
    26: ('Kufiri i auditit', 'U përjashtuan node_modules, .nuxt/.output, raportet e gjeneruara, imazhet binare dhe vlerat sekrete .env. Package/config dhe dokumentet u krahasuan veçmas. Lexim i plotë i burimit first-party nuk garanton mungesë defektesh.'),
    27: ('Nuk u ekzekutuan', 'Clean install/build production, hosted CI, Meta/OpenAI/billing live, Firefox/WebKit, load/restore/pentest/UAT. Asnjë publikim real ose pagesë nuk u krye.'),
    28: ('Si lexohet secili task', '02_BACKLOG: ID në A, status formal L, zbatim Y, kryer Z, mbetur AA, kriter/provë e kërkuar AB, provë aktuale AC, kod/gjetje AD. 15_GJETJET përmban trigger, korrigjim dhe test për çdo defekt.'),
}
for r, (label, value) in updates.items():
    summary.cell(r, 1).value = label
    summary.cell(r, 2).value = value
    for col in [1, 2]:
        summary.cell(r, col)._style = copy(summary.cell(24, col)._style)
        summary.cell(r, col).alignment = Alignment(wrap_text=True, vertical='top')
    summary.row_dimensions[r].height = 88
wb['00_README']['A45'] = 'Rishikim i thelluar: 133 skedarë first-party dhe 185 task-e. 02_BACKLOG: Y zbatimi, Z kryer, AA mbetur, AB pranim/prova të kërkuara, AC prova reale, AD kod/gjetje. 7 In Review, 106 In Progress, 35 Not Started, 2 Blocked, 35 Deferred, 0 Done.'
wb['00_README']['A46'] = 'Done kërkon acceptance, prova dhe gates. Shiko 14_AUDIT_2026-09-07 për metodën/rezultatet dhe 15_GJETJET për 17 gjetjet konkrete. 33 unit/integration dhe 8 E2E kaluan lokalisht; kjo nuk është pranim enterprise.'

issues = wb.create_sheet('15_GJETJET')
issues.append(['ID', 'Prioriteti', 'Gjetja', 'Task-et', 'Kodi', 'Shkaku / sjellja', 'Çfarë mbetet', 'Prova e korrigjimit', 'Niveli i evidencës'])
for f in findings:
    issues.append([f['id'], f['priority'], f['title'], ', '.join(f['tasks']), '\n'.join(f['paths']), f['finding'], f['remaining'], f['verification'], f['evidence']])
for c in issues[1]:
    c.fill = PatternFill('solid', fgColor='173D37')
    c.font = Font(name='Aptos', color='FFFFFF', bold=True)
    c.alignment = Alignment(wrap_text=True, vertical='center')
for r in range(2, issues.max_row + 1):
    issues.row_dimensions[r].height = 155
    for c in issues[r]:
        c.font = Font(name='Aptos', size=11)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.fill = PatternFill('solid', fgColor='F0F6F3' if r % 2 == 0 else 'FFFFFF')
for col, width in enumerate([12, 13, 43, 37, 60, 72, 70, 72, 60], 1):
    issues.column_dimensions[get_column_letter(col)].width = width
issues.row_dimensions[1].height = 35
issues.freeze_panes = 'D2'
issues.auto_filter.ref = issues.dimensions
issues.sheet_view.showGridLines = False
ws.sheet_view.showGridLines = False
wb.active = wb.sheetnames.index('02_BACKLOG')

assert original_formulas == {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == 'f'}
counts = Counter(v['status'] for v in deep.values())
assert counts == Counter(first[t]['status'] for t in first)
with NamedTemporaryFile(dir=ROOT, prefix='backlog-deep-', suffix='.xlsx', delete=False) as handle:
    pending = Path(handle.name)
try:
    wb.save(pending)
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    with ZipFile(pending) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(wb, 1):
        path = f'xl/worksheets/sheet{index}.xml'
        tree = ET.fromstring(files[path])
        for cell in tree.findall(f'.//{{{ns}}}c'):
            if cell.find(f'{{{ns}}}f') is None:
                continue
            value = cache[(sheet.title, cell.attrib['r'])]
            cached = cell.find(f'{{{ns}}}v')
            if cached is None:
                cached = ET.SubElement(cell, f'{{{ns}}}v')
            if isinstance(value, str):
                cell.set('t', 'str')
            else:
                cell.attrib.pop('t', None)
            cached.text = str(value)
        files[path] = ET.tostring(tree, encoding='utf-8', xml_declaration=True)
    with ZipFile(pending, 'w', ZIP_DEFLATED) as archive:
        for name, value in files.items():
            archive.writestr(name, value)
    verified = load_workbook(pending)
    vc = load_workbook(pending, data_only=True)
    result = verified['02_BACKLOG']
    assert result.max_row == 189 and result.max_column == 30
    assert {result.cell(r, 1).value: r for r in range(5, 190)} == rows
    assert sum(result.cell(r, 11).value for r in range(5, 190)) == 1449
    assert all(result.cell(r, c).value for r in range(5, 190) for c in range(25, 31))
    before = load_workbook(BACKUP)
    for sheet in before:
        for row in sheet:
            for cell in row:
                if cell.value == verified[sheet.title][cell.coordinate].value:
                    continue
                allowed = (sheet.title == '02_BACKLOG' and cell.column in [19, 26]) or (
                    sheet.title == summary.title and cell.coordinate in ['A2', 'C6'] +
                    [f'{c}{r}' for r in updates for c in 'AB']) or (
                    sheet.title == '00_README' and cell.coordinate in ['A45', 'A46'])
                assert allowed, (sheet.title, cell.coordinate)
    for key, expected in cache.items():
        actual = vc[key[0]][key[1]].value
        if isinstance(actual, datetime):
            actual = to_excel(actual)
        assert actual == expected, key
    assert vc['01_ROADMAP']['N17'].value == 185
    assert len(cache) == 706
    assert TARGET.read_bytes() == original, 'Workbook changed while editing; preserve concurrent user edits.'
    assert BACKUP.read_bytes() == original
    pending.replace(TARGET)
finally:
    if pending.exists():
        pending.unlink()

write(f'backlog-deep-audit-{DATE}.json', deep)
write('deep-audit-file-inventory.json', inventory)
write('deep-audit-read-coverage.json', coverage)
verification = {
    'date': DATE, 'workbook': NAME, 'tasks': len(deep), 'statuses': dict(counts),
    'source_files_reviewed': len(coverage), 'source_lines': sum(x['lines'] for x in inventory),
    'findings': len(findings), 'original_formulas_and_caches_preserved': len(cache),
    'task_ids_and_estimated_hours_preserved': True, 'hours': 1449,
    'backup': str(BACKUP.relative_to(ROOT)).replace('\\', '/'),
    'before_sha256': sha256(original).hexdigest(), 'after_sha256': sha256(TARGET.read_bytes()).hexdigest(),
    'reviewed_source_changes': source_changes,
    'checks': {'unit_integration': '33 passed', 'e2e_final': '8 passed in two focused runs',
               'initial_live_e2e': '2 passed, 2 failed; DA-17 corrected and rerun 4 passed',
               'lint': 'full exit 0; changed test eslint exit 0', 'app_types': 'exit 0', 'server_types': 'exit 0'},
    'not_executed': ['production build/clean install', 'hosted CI', 'live Meta/OpenAI/billing',
                     'Firefox/WebKit/physical device', 'load/restore/security release/UAT'],
}
write(f'backlog-deep-audit-verification-{DATE}.json', verification)
print(json.dumps(verification, ensure_ascii=True, indent=2))
