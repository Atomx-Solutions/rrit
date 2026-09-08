"""Record the 2026-09-08 partial-task implementation pass in the audited workbook.

The script never promotes a task to Done. It refreshes implementation, remaining
work and executed evidence for the touched partial tasks, preserves formula
expressions and cached values, and keeps a byte-for-byte pre-pass backup.
"""
from collections import Counter
from copy import copy
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile
import json

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils.datetime import to_excel

ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-08"
NAME = "AI_Marketing_Platform_Engineering_Master_Backlog_2_Developers_v1.1.xlsx"
TARGET = ROOT / NAME
BACKUP = ROOT / "docs" / "backups" / NAME.replace(
    ".xlsx", f".before-partial-implementation-{DATE}.xlsx"
)
PROGRESS = ROOT / "docs" / f"backlog-partial-progress-{DATE}.json"


def item(done: str, remaining: str, evidence: str, paths: str):
    return {"done": done, "remaining": remaining, "evidence": evidence, "paths": paths}


UPDATES = {
    "S0-A01": item(
        "ADR-002 është Accepted dhe dokumenton modular monolith Nuxt/Nitro/PostgreSQL, kufijtë e moduleve, worker-in e ndarë dhe deployment-in. ADR-001 është Superseded; README/status/handoff/validation janë sinkronizuar.",
        "Mbeten architecture review checklist, prova e deployment-it të ndarë dhe RG-01. Follow-up-et RLS/runtime-role/observability të ADR-002 nuk janë gates të kaluara.",
        "ADR-002 + kontroll statik i dokumenteve; lint dhe TypeScript kaluan. Nuk ka review sign-off të ekipit ose hosted release gate.",
        "docs/adr/ADR-002-production-modular-monolith.md\ndocs/adr/ADR-001-frontend-demo.md\nREADME.md\ndocs/FRONTEND-STATUS.md\ndocs/BACKEND-HANDOFF.md",
    ),
    "S0-A02": item(
        "API Nitro dhe publisher funksionojnë nga i njëjti artefakt; /api/health/live dhe /api/health/ready ekzistojnë. Readiness kontrollon DB, migrimin 0013, heartbeat-in kur publishing aktivizohet dhe queue age. Kufijtë e moduleve janë në ADR-002.",
        "Duhet container/process smoke me publisher të aktivizuar, deployment i veçantë i worker-it dhe RG-01; Meta App mungon për një start real publishing.",
        "Health live/ready kthyen 200 lokalisht; migrimet 0001–0013 u aplikuan. Nuk u krye container smoke.",
        "apps/web/server/api/health/live.get.ts\napps/web/server/api/health/ready.get.ts\napps/web/server/plugins/publisher.ts\ndocs/adr/ADR-002-production-modular-monolith.md",
    ),
    "S0-A09": item(
        "CI provisionon PostgreSQL, env vetëm testues, aplikon migrimet, ndërton output-in, pret readiness, ekzekuton lint/type/unit/E2E/audit dhe ngarkon artefakte.",
        "Mbeten run hosted, required branch checks, deliberate failing PR dhe gates RLS/secret/container. Workspace nuk ka .git, prandaj merge blocking nuk mund të provohet lokalisht.",
        "Workflow u validua nga lint/lexim dhe konfigurimi lokal ekuivalent kaloi; hosted GitHub Actions nuk u ekzekutua.",
        ".github/workflows/frontend.yml\napps/web/server/api/health/ready.get.ts\n.nvmrc",
    ),
    "S0-B01": item(
        "Nuxt 4.5.2/Nuxt UI 4.11.0 me TypeScript strict; .nvmrc tani fikson Node 24. Lint dhe app/server typecheck kalojnë dhe sekretet mbeten server-side.",
        "Duhet npm ci + production build nga checkout i pastër me Node 24 dhe RG-01. Makina e këtij auditimi po përdor Node 26.5.0.",
        "ESLint exit 0; vue-tsc app exit 0; tsc server exit 0. Clean Node 24 build nuk u ekzekutua.",
        ".nvmrc\napps/web/package.json\napps/web/nuxt.config.ts\nhq-partials-lint-2.log",
    ),
    "S0-B04": item(
        "Login, signup dhe recovery kanë capability të ndara; trial/SMTP i munguar nuk bllokon login-in. Regjistrimi lokal, logout/login, localhost/127.0.0.1 dhe foreign-origin rejection kaluan në browser.",
        "Mbeten verify/reset/expired/error me SMTP real në staging, public-origin verification dhe RG-03.",
        "47/47 unit/integration; 8/8 E2E desktop/mobile. Testet e auth capability mbulojnë DB/trial/SMTP degradation.",
        "apps/web/server/utils/auth-configuration.ts\napps/web/server/utils/auth.ts\napps/web/server/api/capabilities.get.ts\napps/web/app/components/AuthSurface.vue\napps/web/tests/unit/auth-configuration.test.ts\napps/web/tests/e2e/local-registration.spec.ts",
    ),
    "S0-B07": item(
        "Pipeline-i përmban Node 24, PostgreSQL, migrime, lint, typecheck, unit/integration, audit, build, bundle budget, readiness, E2E dhe artifacts.",
        "Mbeten run hosted, reproducibility nga checkout i pastër, required-check enforcement dhe deliberate failing check.",
        "Të gjitha kontrollet ekuivalente përveç clean build/hosted CI kaluan lokalisht; kjo nuk provon merge blocking.",
        ".github/workflows/frontend.yml\n.nvmrc\ndocs/VALIDATION.md",
    ),
    "S0-B08": item(
        "Playwright ekzekuton desktop/mobile me trace dhe screenshot në dështim. Tetë teste të kombinuara kaluan me exit 0; rate limit i signup mbetet aktiv dhe lejon suitën deterministe.",
        "Mbeten hosted CI run, cleanup/isolim i provuar i fixture-ve dhe RG-14. Emrat aksesues të pagination-it mbeten anglisht nga primitive-i.",
        "8/8 E2E në Chromium desktop/mobile, përfshirë regjistrim, private tenant, pagination dhe optional Meta.",
        "apps/web/playwright.config.ts\napps/web/tests/e2e/live.spec.ts\napps/web/tests/e2e/local-registration.spec.ts\napps/web/tests/e2e/optional-meta.spec.ts",
    ),
    "S2-A07": item(
        "Migrimi 0008 ruan source_account_id; uniciteti, query-t, ndërrimi i faqes dhe konteksti AI filtrohen sipas llogarisë aktuale. Provider ID i njëjtë në dy faqe nuk përzihet.",
        "Mbeten durable incremental worker, checkpoint/cursor i ruajtur, crash resume, revoked-auth taxonomy, inference dhe test live Meta.",
        "Integration test provon dy source accounts dhe duplicate brenda të njëjtit source. Nuk u krye ndërrim me dy faqe reale.",
        "apps/web/database/migrations/0008_provider_reliability.sql\napps/web/server/api/meta/sync.post.ts\napps/web/server/api/meta/posts.get.ts\napps/web/server/api/meta/select.post.ts\napps/web/server/api/ai/captions.post.ts",
    ),
    "S2-B06": item(
        "Zgjedhja e faqes ruan target-in aktual dhe ndalon kalimin te faqe tjetër kur një rezultat uncertain i përket burimit të vjetër. Rilidhja lejohet vetëm me target-in origjinal.",
        "Mbeten capability matrix për missing scopes/expired/revoked/unsupported, permission reauth dhe E2E me Meta staging.",
        "Unit/integration provon target page/Instagram të njëjtë dhe refuzon target tjetër; prova live provider mungon.",
        "apps/web/server/api/meta/select.post.ts\napps/web/server/api/meta/callback.get.ts\napps/web/server/utils/publishing.ts",
    ),
    "S2-B07": item(
        "Historiku filtrohet sipas source account aktual dhe pastrohet gjatë ndërrimit të faqes; UI ruan pagination/search/last sync.",
        "Mbeten job/progress/history durable, crash resume dhe terminal-auth/retry fixtures.",
        "DB integration për source isolation kaloi; nuk ka provider sync fault E2E.",
        "apps/web/server/api/meta/posts.get.ts\napps/web/server/api/meta/sync.post.ts\napps/web/database/migrations/0008_provider_reliability.sql",
    ),
    "S3-A05": item(
        "AI style history përdor vetëm postimet e source account-it të lidhur aktualisht; profil/facts verified mbeten konteksti bazë.",
        "Mbetet evidence pack immutable me IDs, hashes, voice/offer validity, size limits dhe inclusion/exclusion suite.",
        "Source separation ka DB test; evidence compiler i plotë nuk ekziston.",
        "apps/web/server/api/ai/captions.post.ts\napps/web/database/migrations/0008_provider_reliability.sql",
    ),
    "S3-A07": item(
        "Migrimet 0009/0010 ndajnë contentVersion nga workflow version, ruajnë snapshot të plotë immutable me actor/source/fingerprint dhe lidhin validation, approval dhe publication job me të njëjtin version/tenant. Editimi krijon version të ri.",
        "Mbeten concurrency test për dy writer-a në endpoint, verifikim i published-version me provider fixture dhe RG-06. Historia caption para 0009 nuk mund të rikonstruktohet si snapshot i plotë dhe shënohet legacy.",
        "Integration test provon immutability trigger, approval që nuk kalon te versioni i ri dhe FK që refuzon tenant tjetër; 47/47 kaluan.",
        "apps/web/database/migrations/0009_immutable_content_versions.sql\napps/web/database/migrations/0010_content_tenant_integrity.sql\napps/web/server/api/workspaces/[id].put.ts\napps/web/server/utils/post-validation.ts\napps/web/tests/integration/server.test.ts",
    ),
    "S3-B06": item(
        "Backend-i ruan snapshot të plotë immutable për çdo edit të ri; historia caption ekzistuese nuk muton versionin e vjetër.",
        "Mbeten UI diff/side-by-side për tekst/foto/kanal, provenance e dukshme, zgjedhje varianti dhe E2E mobile comparison.",
        "DB immutability test kaloi; component/E2E comparison nuk ekziston.",
        "apps/web/database/migrations/0009_immutable_content_versions.sql\napps/web/app/components/ContentEditor.vue",
    ),
    "S4-A01": item(
        "State machine, fresh validation dhe approval records immutable tani referojnë contentVersion/actor/fingerprint. Editi krijon version të ri pa miratimin e vjetër; re-open regjistrohet.",
        "Mbeten roli/capability Approver, policy override me reason, role matrix dhe RG-03/RG-07.",
        "Integration test provon version-bound approval; E2E optional Meta provon review→approve për dy kanale.",
        "apps/web/server/api/workspaces/[id].put.ts\napps/web/database/migrations/0009_immutable_content_versions.sql\napps/web/server/utils/post-validation.ts",
    ),
    "S4-A03": item(
        "Dispatcher-i përdor SKIP LOCKED, fairness sipas punës due më të vjetër, batch 1–25 dhe poll 1–60s. Dy worker-a paralelë marrin vetëm një job dhe readiness raporton pendingJobs/oldestDueSeconds/heartbeat.",
        "Mbeten burst/load test, outbox/dispatch proof, procesi i veçantë në deployment dhe alert/SLO për queue age.",
        "Fixture-i PostgreSQL me dy worker-a paralelë kaloi me vetëm një shkrim provider; health readiness ktheu 200. Nuk ka burst benchmark.",
        "apps/web/server/utils/publishing.ts\napps/web/server/plugins/publisher.ts\napps/web/server/api/health/ready.get.ts",
    ),
    "S4-A04": item(
        "Instagram container_start ruhet veçmas nga schedule creation; timeout matet nga fillimi real. Publication job lidhet me contentVersion immutable. Provider fixture ekzekuton worker-in dhe verifikon vetëm një Facebook write.",
        "Mbeten Meta staging publish me llogari të autorizuar, public HTTPS/media expiry live dhe matrica e plotë për Instagram/provider errors.",
        "Fixture-i worker/provider dhe testi 59/61 minuta kaluan. Asnjë postim real nuk u dërgua.",
        "apps/web/database/migrations/0008_provider_reliability.sql\napps/web/server/utils/publishing.ts",
    ),
    "S4-A05": item(
        "Timeout-i i container-it u korrigjua dhe uncertain/failed nuk riprovohen verbërisht; failure codes ruhen.",
        "Mbeten Retry-After, exponential backoff/jitter, attempt limits dhe taxonomy fixtures për 429/5xx/network/auth/media.",
        "Boundary unit test kaloi; provider fault matrix mungon.",
        "apps/web/server/utils/publishing.ts\napps/web/server/utils/meta.ts",
    ),
    "S4-A06": item(
        "Reconciler-i rilidh vetëm target-in origjinal, verifikon provider ID ose një përputhje unike caption/time. Fault fixture simulon timeout pas side-effect-it, kalon në uncertain dhe e mbyll published pa ridërgim.",
        "Kriteri lokal i simulimit është përmbushur. Mbeten staging proof me një remote post dhe procedura operative për zero/shumë përputhje; RG-07 nuk është mbyllur.",
        "Integration test ekzekuton timeout-after-provider-write, verifikon një remote write logjik dhe reconciliation pa POST të dytë; live Meta nuk u përdor.",
        "apps/web/server/api/posts/[id]/reconcile.post.ts\napps/web/server/utils/publishing.ts\napps/web/server/api/meta/select.post.ts",
    ),
    "S4-A07": item(
        "Unique job guard, SKIP LOCKED dhe worker fixture provojnë se dy delivery/worker-a paralelë japin një publikim logjik; provider receipt ruhet përpara verifikimit dhe uncertain nuk ridërgohet verbërisht.",
        "Mbetet publication_attempts ledger immutable me request hash dhe histori të renditur për çdo tentim, plus prova HTTP→queue→provider në staging.",
        "Dy runPublishingBatch paralelë prodhuan vetëm një provider POST; timeout-after-write u pajtua pa POST tjetër.",
        "apps/web/database/migrations/0004_trials_and_social.sql\napps/web/server/utils/publishing.ts\napps/web/tests/integration/server.test.ts",
    ),
    "S4-A09": item(
        "Suite mbulon success me dy worker-a, duplicate claim, timeout-after-write, safe reconciliation, container timeout, target binding, signatures, validation scopes dhe manual/automatic guards.",
        "Mbeten 429/Retry-After, 5xx, auth expiry, media reject, cancel race, process crash dhe provider fixtures në hosted CI/staging.",
        "47/47 unit/integration kaluan; success/duplicate/ambiguous outcome tani ekzekutojnë worker-in real me fetch provider fixture.",
        "apps/web/tests/integration/server.test.ts\napps/web/server/utils/publishing.ts",
    ),
    "S4-B01": item(
        "Reviewer-i sheh përmbajtjen aktuale; approval record i pandryshueshëm lidhet me snapshot-in e plotë dhe stale workspace revision kthen konflikt. E2E approve kalon për Facebook/Instagram.",
        "Mbeten Approver role/unauthorized matrix, comments/reject reason, UI që hap snapshot historik dhe RG-03/RG-14.",
        "DB version/tenant tests + 8/8 browser tests kaluan.",
        "apps/web/app/components/ContentEditor.vue\napps/web/server/api/workspaces/[id].put.ts\napps/web/database/migrations/0009_immutable_content_versions.sql",
    ),
    "S4-B04": item(
        "Failed ka rikthim në draft ose ri-miratim; uncertain shfaq paralajmërim, link për lidhjen dhe kontroll të sigurt në Meta. Provider/failure ID nuk ekspozon token.",
        "Mbeten attempt timeline/support correlation ID, remediation sipas çdo kodi dhe E2E provider states.",
        "State tests dhe TypeScript/UI checks kaluan; Meta attempt E2E mungon.",
        "apps/web/app/components/ContentEditor.vue\napps/web/server/api/posts/[id]/publication.get.ts\napps/web/server/api/posts/[id]/reconcile.post.ts",
    ),
    "S4-B06": item(
        "UI ofron failed recovery, uncertain reconciliation dhe reconnect link. Serveri lejon vetëm target-in origjinal dhe nuk ridërgon pa konfirmim unik.",
        "Mbetet error-code UX matrix për Retry-After/quota/media/auth dhe provider fault E2E.",
        "Unit/integration target/match tests + lint/types kaluan; live provider remediation mungon.",
        "apps/web/app/components/ContentEditor.vue\napps/web/server/api/meta/callback.get.ts\napps/web/server/api/meta/select.post.ts\napps/web/server/utils/publishing.ts",
    ),
    "S4-B07": item(
        "DB ruan content snapshots, approval actor/time/version, validations dhe publication jobs të lidhura me contentVersion.",
        "Mbetet endpoint/UI timeline i unifikuar generation/approval/schedule/publication dhe E2E që kontrollon renditjen/redaction.",
        "Integriteti dhe immutability kaluan në integration; timeline UI nuk është ndërtuar.",
        "apps/web/database/migrations/0009_immutable_content_versions.sql\napps/web/database/migrations/0010_content_tenant_integrity.sql",
    ),
    "S1-A07": item(
        "Trial-i dhe paid entitlements tani përdorin feature codes dhe limite të tipizuara sipas planit/versionit në PostgreSQL. Skema refuzon katalog të pavlefshëm, trial-i është 72 orë dhe nuk zgjatet në retry.",
        "Mbeten subscription/provider lifecycle, downgrade rules, device abuse hardening dhe prova e integrimit komercial end-to-end; RG-08 mbetet i hapur.",
        "Integration tests lexojnë plan/version/limits nga DB, validojnë katalogun EUR dhe provojnë trial/paid/expired/blocked/revoked. 47/47 kaluan.",
        "apps/web/database/migrations/0011_typed_plan_catalog.sql\napps/web/shared/schemas/catalog.ts\napps/web/server/utils/trials.ts\napps/web/tests/integration/server.test.ts",
    ),
    "S2-B03": item(
        "Asset-i ruan metadata dhe të drejtat; UI lejon tërheqjen e konfirmuar të të drejtave. Serveri anulon job-et pending, kthen postimet e papublikuara në draft, i bllokon nga ripërdorimi dhe ruan audit/approval history. Preview ruan raportin natyral të medias.",
        "Mbeten variantet/crop processing, object storage production dhe suite më e gjerë metadata/delete/concurrency; publikimet tashmë in-flight kërkojnë reconciliation operacional.",
        "E2E optional-Meta provon tërheqjen e të drejtave; local validation e refuzon asset-in pas tërheqjes. 8/8 browser dhe 47/47 unit/integration kaluan.",
        "apps/web/server/api/media/[id]/rights.post.ts\napps/web/app/components/MediaGallery.vue\napps/web/app/components/ChannelPostPreview.vue\napps/web/tests/e2e/optional-meta.spec.ts",
    ),
    "S3-B02": item(
        "Editor-i ruan draftin në server, trajton conflict/version, përdor media private dhe preview Facebook/Instagram me raport natyral të njëjtë me asset-in. Kombinimet kanal/format lexohen nga registry i përbashkët.",
        "Mbeten fault/browser matrix më e plotë, composer me më shumë se një asset dhe pranimi RG-06/RG-14.",
        "Registry unit test, app/server types dhe approve/calendar E2E desktop/mobile kaluan.",
        "apps/web/app/components/ContentEditor.vue\napps/web/app/components/ChannelPostPreview.vue\napps/web/shared/utils/content-formats.ts\napps/web/tests/unit/workspace.test.ts",
    ),
    "S5-A06": item(
        "Endpoint-et server kërkojnë feature entitlement dhe kufijtë posts/media/AI lexohen nga plani aktiv. Denial-i kthen kod të qëndrueshëm; /api/usage kthen current/limit/state për çdo quota key.",
        "Mbeten downgrade state machine, members enforcement dhe matrica e parametrizuar për çdo plan/endpoint; RG-08 nuk është kaluar.",
        "DB tests provojnë planin e tipizuar, expiry dhe feature denial; public plan catalog dhe usage schema kalojnë typecheck.",
        "apps/web/database/migrations/0011_typed_plan_catalog.sql\napps/web/server/utils/trials.ts\napps/web/server/api/usage.get.ts\napps/web/server/api/workspaces/[id].put.ts",
    ),
    "S5-A07": item(
        "AI ka ledger monetar reserve/commit/release me rezervim unik për generation, input/output usage dhe kosto të llogaritur nga çmime të konfiguruara. Dështimi liron rezervimin dhe rimburson quota count vetëm një herë.",
        "Mbeten concurrency burst/fault tests për endpoint-in me provider, reconciler për rezervime të ngecura dhe policy e aprovuar për kreditë/soft budget.",
        "Integration test provon llogaritjen e kostos dhe vetëm një rezervim për generation; constraint-et DB dhe lint/types kaluan.",
        "apps/web/database/migrations/0013_ai_cost_ledger.sql\napps/web/server/utils/ai-costs.ts\napps/web/server/api/ai/captions.post.ts\napps/web/tests/integration/server.test.ts",
    ),
    "S5-B03": item(
        "Landing pricing lexon katalogun e tipizuar nga PostgreSQL: EUR, version, entitlements, limits dhe purchasable. Nuk ka matricë komerciale runtime të duplikuar dhe plani unavailable nuk hap checkout.",
        "Mbeten miratimi i çmimeve, provider checkout/portal dhe contract test i subscription→entitlement; RG-08/RG-14 mbeten të hapura.",
        "Plan catalog schema dhe DB integration test kaluan; /api/plans ktheu katalogun EUR.",
        "apps/web/database/migrations/0011_typed_plan_catalog.sql\napps/web/shared/schemas/catalog.ts\napps/web/server/api/plans.get.ts\napps/web/app/pages/pricing.vue",
    ),
    "S5-B05": item(
        "Billing UI lexon /api/usage dhe shfaq current/limit, përqindjen dhe gjendjet available/warning/exhausted për postime, AI, storage, llogari sociale dhe anëtarë.",
        "Mbeten refresh E2E pas gjenerimit/ngarkimit, reserved vs consumed visibility dhe copy/CTA e aprovuar për kufijtë 80/100%.",
        "Usage snapshot DB test dhe TypeScript kaluan; endpoint-i llogarit gjendjen në server.",
        "apps/web/server/api/usage.get.ts\napps/web/server/utils/trials.ts\napps/web/app/pages/app/[business]/billing.vue",
    ),
    "S5-B06": item(
        "Feature gating mbështetet te feature codes/limits të tipizuara dhe backend denial; UI e usage/capabilities nuk kontrollon emra planesh të shpërndarë.",
        "Mbeten composable i unifikuar për upgrade context, stale-snapshot E2E dhe matrica e çdo denial-i/CTA-je.",
        "Catalog/entitlement integration tests, lint dhe app/server types kaluan.",
        "apps/web/shared/schemas/catalog.ts\napps/web/server/utils/trials.ts\napps/web/server/api/capabilities.get.ts\napps/web/server/api/usage.get.ts",
    ),
    "S6-A13": item(
        "AI cost ledger rezervon buxhetin para provider call, commit-on-success, release-on-failure dhe kill switch përmes konfigurimit. Çmimet dhe hard daily budget janë server-side dhe readiness e capability dështon mbyllur kur mungojnë.",
        "Mbeten synthetic concurrent load, anomaly alerts/operator recovery, stuck-reservation reconciliation dhe envelope komercial i aprovuar; RG-18 mbetet i hapur.",
        "Cost calculation/unique reservation integration test, DB constraints, lint dhe server types kaluan.",
        "apps/web/database/migrations/0013_ai_cost_ledger.sql\napps/web/server/utils/ai-costs.ts\napps/web/server/api/ai/captions.post.ts\n.env.example",
    ),
    "S7-A03": item(
        "ContentFormat/ChannelCapability registry qendror përcakton mbështetjen dhe mediaFit për Facebook/Instagram. Validimi dhe preview përdorin të njëjtin burim; carousel dështon para remote publish.",
        "Mbeten reel/story/carousel adapters, account-level capability contract dhe provider contract tests për secilin format.",
        "Unit test provon registry-n static/carousel; local checks dhe app/server typecheck kaluan.",
        "apps/web/shared/utils/content-formats.ts\napps/web/server/utils/post-validation.ts\napps/web/app/components/ChannelPostPreview.vue\napps/web/tests/unit/workspace.test.ts",
    ),
    "S7-A05": item(
        "Retention policy ruhet në PostgreSQL dhe maintenance job pastron të dhëna të skaduara me advisory lock. Tërheqja e të drejtave të medias është e enforce-ueshme, anulon përdorimin e ardhshëm dhe ruan historikun/auditin.",
        "Mbeten provider-authenticated deletion callback, organization purge resumable, retention classes/holds dhe legal owner approval; RG-09/RG-16 mbeten të hapura.",
        "Retention integration test dhe rights-withdrawal browser flow kaluan; 0012/0013 janë aplikuar lokalisht.",
        "apps/web/database/migrations/0012_retention_policy.sql\napps/web/server/utils/retention.ts\napps/web/server/plugins/maintenance.ts\napps/web/server/api/media/[id]/rights.post.ts",
    ),
    "S8-A10": item(
        "Retention policy/job dhe rights-withdrawal lifecycle tani kanë implementim dhe prova lokale; dokumentet e platformës dallojnë qartë kufirin lokal nga gates production.",
        "Mbeten export/purge runbooks të provuara, deletion callback, inventory/subprocessors/DPA dhe miratim ligjor i pronarit; RG-09 nuk është mbyllur.",
        "Retention DB test, rights E2E dhe dokumentet e sinkronizuara kaluan auditin lokal.",
        "apps/web/database/migrations/0012_retention_policy.sql\napps/web/server/utils/retention.ts\napps/web/server/api/media/[id]/rights.post.ts\ndocs/BACKEND-HANDOFF.md",
    ),
    "S5-A09": item(
        "U shtuan index-i i source history, batch/fairness dhe queue age në readiness. Query-t kyçe kanë indekse të versionuara.",
        "Mbeten server-side list APIs, 500+/1000 fixture, EXPLAIN/k6 report, p95/memory evidence dhe RG-11.",
        "Kontrollet funksionale kalojnë; nuk u mat staging load.",
        "apps/web/database/migrations/0008_provider_reliability.sql\napps/web/server/api/health/ready.get.ts\napps/web/server/utils/publishing.ts",
    ),
    "S6-B05": item(
        "Auth degradation është ndarë sipas veprimit; failed/uncertain kanë rrugë rikuperimi pa humbur draftin ose pa dyfishuar publikimin.",
        "Mbetet network throttle/offline/intermittent/provider fault E2E dhe RG-14.",
        "Auth unit tests, state tests dhe 8 E2E kaluan në mjedis normal; fault injection browser mungon.",
        "apps/web/server/utils/auth-configuration.ts\napps/web/app/components/AuthSurface.vue\napps/web/app/components/ContentEditor.vue",
    ),
    "S7-B03": item(
        "Approval/validation lidhen me snapshot immutable; postimi failed mund të kontrollohet dhe miratohet sërish. BLOCK vazhdon të ndalë approve/schedule.",
        "Mbeten factuality/claim provenance, versioned policy, override capability/reason/audit dhe RG-16/RG-17.",
        "State/DB tests kaluan; override/provenance acceptance nuk është kryer.",
        "apps/web/server/utils/post-validation.ts\napps/web/server/api/workspaces/[id].put.ts\napps/web/app/components/ContentEditor.vue",
    ),
    "S7-B08": item(
        "Mobile Chromium 390px kalon regjistrim, approval/calendar pa Meta dhe pagination; failed/uncertain veprimet përdorin Nuxt UI controls pa hover-only dependency.",
        "Mbeten 360px, pajisje fizike/Safari, reject/deep-link/reschedule/reconnect scripted flow dhe touch-target review.",
        "8/8 E2E desktop/mobile kaluan; real-device spot nuk u krye.",
        "apps/web/app/components/ContentEditor.vue\napps/web/tests/e2e/optional-meta.spec.ts\napps/web/playwright.config.ts",
    ),
    "S8-A11": item(
        "Përveç version/tenant guards dhe container timeout, provider fixture tani ekzekuton dy worker-a, success, duplicate claim dhe timeout-after-write me reconciliation pa ridërgim.",
        "Mbetet staged campaign për 429/5xx/auth/media/process crash, AI provider dhe billing me alerts/SLO; nuk ka fault campaign të jashtëm.",
        "47/47 unit/integration kaluan, përfshirë worker/provider ambiguous-outcome fixture; billing dhe provider live nuk u thirrën.",
        "apps/web/tests/integration/server.test.ts\napps/web/server/utils/publishing.ts\napps/web/database/migrations/0009_immutable_content_versions.sql",
    ),
    "S8-B01": item(
        "Tetë teste kritike kaluan në desktop/mobile Chromium me DB reale lokale dhe migrimet deri në 0013.",
        "Mbeten Firefox/WebKit, hosted CI/staging dhe flows reale generation/publishing/billing.",
        "8/8 Playwright exit 0 në 1.1 minuta; vetëm Chromium.",
        "apps/web/tests/e2e\napps/web/playwright.config.ts\ndocs/VALIDATION.md",
    ),
    "S8-B12": item(
        "Fresh accounts provojnë signup lokal, empty workspace, login pas logout, pa Meta, approval/calendar, invalid content dhe capability-specific auth degradation.",
        "Mbeten revoked account, quota exhausted/provider degraded browser suite dhe UAT me përdorues real.",
        "8/8 E2E + auth degradation unit tests kaluan; UAT nuk u krye.",
        "apps/web/tests/e2e/local-registration.spec.ts\napps/web/tests/e2e/optional-meta.spec.ts\napps/web/tests/unit/auth-configuration.test.ts",
    ),
    "S8-B13": item(
        "Dokumentet/copy ndjekin backend-in real; help/pricing/settings lexojnë capability e verifikimit, Meta është opsionale dhe nuk ka runtime demo. Fresh account ruan progres real.",
        "Mbeten checklist/funnel telemetry, UAT me biznese reale dhe RG-17.",
        "Content review statik, lint/types dhe fresh-account E2E kaluan. Product telemetry/UAT mungojnë.",
        "apps/web/app/pages/help.vue\napps/web/app/pages/pricing.vue\napps/web/app/components/CopilotPanel.vue\ndocs/FRONTEND-STATUS.md\ndocs/PRODUCT-SCOPE.md",
    ),
}

FINDINGS = {
    "DA-01": ("Resolved locally", "container_started_at + 59/61 minute boundary test; live Meta remains a task gate."),
    "DA-02": ("Partially addressed", "Bounded batch, oldest-due fairness, configurable poll and queue-age readiness; multi-worker load proof remains."),
    "DA-03": ("Partially addressed", "Failed can return to draft or re-approval; full provider error-code E2E remains."),
    "DA-04": ("Resolved locally", "Timeout-after-side-effect provider fixture reconciles one exact result without a second write; live Meta remains a release gate."),
    "DA-05": ("Resolved locally", "source_account_id migration/query/AI filtering and two-source DB test pass; live two-page test remains a release gate."),
    "DA-06": ("Resolved locally", "Login/signup/recovery readiness split and degradation tests pass."),
    "DA-07": ("Resolved locally", "Full immutable snapshots and version-bound validation/approval/attempt records with tenant FKs pass DB tests."),
    "DA-08": ("Partially addressed", "CI now provisions PostgreSQL/env/migrations/readiness; hosted required-check proof remains."),
    "DA-09": ("Open", "Server-side list/command APIs and load evidence remain."),
    "DA-10": ("Open", "Restorable server export with media/manifest remains."),
    "DA-11": ("Partially addressed", "Versioned retention batch and enforceable rights withdrawal exist; purge/deletion callbacks and legal approval remain."),
    "DA-12": ("Partially addressed", "AI reserve/commit/release cost ledger and hard configured budget exist; live concurrency/replay and anomaly operations remain."),
    "DA-13": ("Resolved locally", "ADR/README/status/handoff/validation and capability-aware UI copy updated; public content sign-off remains."),
    "DA-14": ("Resolved locally", "Provider fixture executes two workers, one write and timeout-after-write reconciliation; staging publish remains a release gate."),
    "DA-15": ("Partially addressed", "Typed versioned catalog, limits, enforcement and usage snapshot exist; subscription lifecycle/checkout and commercial approval remain."),
    "DA-16": ("Resolved locally", "Preview preserves the real media ratio and validation/preview share one channel-format registry; real-provider visual proof remains a task gate."),
    "DA-17": ("Resolved locally", "Pagination locator fixed and combined E2E passes 8/8; hosted CI/localized primitive labels remain."),
}

original = TARGET.read_bytes()
BACKUP.parent.mkdir(parents=True, exist_ok=True)
if not BACKUP.exists():
    BACKUP.write_bytes(original)

wb = load_workbook(TARGET)
values = load_workbook(TARGET, data_only=True)
formulas = {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == "f"}
cache = {}
for key in formulas:
    value = values[key[0]][key[1]].value
    if value is None:
        raise AssertionError(f"Missing cached value for {key}")
    cache[key] = to_excel(value) if isinstance(value, datetime) else value

ws = wb["02_BACKLOG"]
rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1)}
if not set(UPDATES) <= set(rows):
    raise AssertionError("Task ID mismatch")
before_statuses = Counter(ws.cell(r, 12).value for r in range(5, ws.max_row + 1))
for task_id, update in UPDATES.items():
    row = rows[task_id]
    if ws.cell(row, 12).value not in {"In Progress", "In Review"}:
        raise AssertionError(f"Refusing to edit non-partial task {task_id}")
    ws.cell(row, 19).value = f"docs/backlog-partial-progress-{DATE}.json#{task_id}"
    ws.cell(row, 19).hyperlink = f"docs/backlog-partial-progress-{DATE}.json"
    ws.cell(row, 26).value = update["done"]
    ws.cell(row, 27).value = update["remaining"]
    ws.cell(row, 29).value = update["evidence"]
    ws.cell(row, 30).value = update["paths"]
    for col in range(26, 31):
        ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 180

issues = wb["15_GJETJET"]
issues.cell(1, 10).value = "Statusi pas implementimit"
issues.cell(1, 11).value = "Evidenca e re / kufiri"
for col in (10, 11):
    issues.cell(1, col)._style = copy(issues.cell(1, 9)._style)
    issues.cell(1, col).alignment = Alignment(wrap_text=True, vertical="center")
for row in range(2, issues.max_row + 1):
    finding_id = issues.cell(row, 1).value
    status, proof = FINDINGS[finding_id]
    issues.cell(row, 10).value = status
    issues.cell(row, 11).value = proof
    for col in (10, 11):
        issues.cell(row, col)._style = copy(issues.cell(row, 9)._style)
        issues.cell(row, col).alignment = Alignment(wrap_text=True, vertical="top")
issues.column_dimensions["J"].width = 22
issues.column_dimensions["K"].width = 74
issues.auto_filter.ref = f"A1:K{issues.max_row}"

audit = wb["14_AUDIT_2026-09-07"]
audit["A30"] = "Progresi i implementimit të task-eve partial — 2026-09-08"
audit["B30"] = (
    f"{len(UPDATES)} task-e partial u rifreskuan; 0 u promovuan artificialisht në Done. "
    "Migrimet 0001–0013, 47/47 unit/integration, 8/8 E2E, lint dhe app/server types kaluan. "
    "Shiko 15_GJETJET kolonat J–K dhe backlog kolonat Z/AA/AC/AD."
)
for col in (1, 2):
    audit.cell(30, col)._style = copy(audit.cell(29, col)._style)
    audit.cell(30, col).alignment = Alignment(wrap_text=True, vertical="top")
audit.row_dimensions[30].height = 90
wb["00_README"]["A47"] = (
    f"Progress 2026-09-08: {len(UPDATES)} task-e partial u përditësuan me implementim dhe prova të reja; "
    "0 u quajtën Done pa gates. 47/47 unit/integration dhe 8/8 E2E kaluan, bashkë me lint/types. "
    "15_GJETJET J–K tregon statusin aktual të 17 gjetjeve."
)

after_statuses = Counter(ws.cell(r, 12).value for r in range(5, ws.max_row + 1))
if before_statuses != after_statuses or after_statuses.get("Done", 0) != 0:
    raise AssertionError("Formal statuses changed unexpectedly")
if formulas != {(s.title, c.coordinate): c.value for s in wb for row in s for c in row if c.data_type == "f"}:
    raise AssertionError("Formula expressions changed")

with NamedTemporaryFile(dir=ROOT, prefix="backlog-progress-", suffix=".xlsx", delete=False) as handle:
    pending = Path(handle.name)
try:
    wb.save(pending)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(pending) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    for index, sheet in enumerate(wb, 1):
        path = f"xl/worksheets/sheet{index}.xml"
        tree = ET.fromstring(files[path])
        for cell in tree.findall(f".//{{{namespace}}}c"):
            if cell.find(f"{{{namespace}}}f") is None:
                continue
            cached_value = cache[(sheet.title, cell.attrib["r"])]
            cached = cell.find(f"{{{namespace}}}v")
            if cached is None:
                cached = ET.SubElement(cell, f"{{{namespace}}}v")
            if isinstance(cached_value, str):
                cell.set("t", "str")
            else:
                cell.attrib.pop("t", None)
            cached.text = str(cached_value)
        files[path] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with ZipFile(pending, "w", ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    check = load_workbook(pending)
    check_values = load_workbook(pending, data_only=True)
    if len(formulas) != 706:
        raise AssertionError("Unexpected formula count")
    for key, expected in cache.items():
        actual = check_values[key[0]][key[1]].value
        if isinstance(actual, datetime):
            actual = to_excel(actual)
        if actual != expected:
            raise AssertionError(f"Formula cache changed: {key}")
    if TARGET.read_bytes() != original:
        raise AssertionError("Workbook changed concurrently")
    pending.replace(TARGET)
finally:
    if pending.exists():
        pending.unlink()

payload = {
    "date": DATE,
    "generated_at": datetime.now(timezone.utc).isoformat(),
    "workbook": NAME,
    "partial_tasks_updated": len(UPDATES),
    "formal_statuses_unchanged": dict(after_statuses),
    "done_promotions": 0,
    "checks": {
        "migrations": "0001-0013 applied",
        "unit_integration": "47/47 passed",
        "e2e": "8/8 passed, Chromium desktop/mobile",
        "lint": "exit 0",
        "app_types": "exit 0",
        "server_types": "exit 0",
        "readiness": "200; database/migrations ready, publisher disabled locally",
    },
    "tasks": UPDATES,
    "findings": {key: {"status": value[0], "evidence": value[1]} for key, value in FINDINGS.items()},
    "backup": str(BACKUP.relative_to(ROOT)).replace("\\", "/"),
    "before_sha256": sha256(original).hexdigest(),
    "after_sha256": sha256(TARGET.read_bytes()).hexdigest(),
    "formula_expressions_and_caches_preserved": len(formulas),
}
PROGRESS.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({key: payload[key] for key in ("partial_tasks_updated", "formal_statuses_unchanged", "done_promotions", "checks", "backup", "formula_expressions_and_caches_preserved")}, ensure_ascii=True, indent=2))
